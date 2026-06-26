import asyncio
import io
import json
import re
from typing import Any

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from starlette.concurrency import run_in_threadpool

from core.utils.email_verify import verify_email_legitimacy
from core.utils.lead_summary import build_plain_lead_summary
from core.utils.phone_company_match import PhoneCompanyMatcher
from core.utils.phone_verify import verify_phone_legitimacy


router = APIRouter()

EMAIL_PATTERN = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.IGNORECASE)
MAX_UPLOAD_BYTES = 8 * 1024 * 1024
DEFAULT_STATUS_HEADER = "email_verification_status"
DEFAULT_SUMMARY_HEADER = "Lead Summary"
EMPTY_EMAIL_VALUES = {"", "-", "--", "—", "n/a", "na", "none", "null", "not found"}
EMAIL_HEADER_ALIASES = {
    "email",
    "emailaddress",
    "emailid",
    "contactemail",
    "contactemailaddress",
    "personemail",
    "leademail",
    "companyemail",
    "workemail",
    "businessemail",
    "e-mail",
    "mail",
}
PHONE_HEADER_ALIASES = {
    "phone",
    "phonenumber",
    "contactphone",
    "contactnumber",
    "mobile",
    "mobilenumber",
    "tel",
    "telephone",
    "cell",
    "cellphone",
    "workphone",
    "businessphone",
}
WEBSITE_HEADER_ALIASES = {
    "website",
    "companywebsite",
    "url",
    "companyurl",
    "domain",
    "web",
    "site",
    "officialwebsite",
}
SUMMARY_HEADER_ALIASES = {"leadsummary", "businesssummary", "plainenglishsummary", "whattheydo", "summary"}
PHONE_COMPANY_MATCH_HEADERS = {
    "status": "phone_company_match_status",
    "source": "phone_company_match_source",
    "confidence": "phone_company_match_confidence",
    "matched_phone": "phone_company_matched_number",
    "scraped_phones": "phone_numbers_found_on_website",
}
HEADER_TO_LEAD_KEY = {
    "companyname": "company_name",
    "company": "company_name",
    "business": "company_name",
    "specialty": "specialty",
    "industry": "industry",
    "category": "industry",
    "contactname": "contact_person_name",
    "personname": "contact_person_name",
    "founder": "contact_person_name",
    "owner": "contact_person_name",
    "email": "contact_email",
    "contactemail": "contact_email",
    "workemail": "contact_email",
    "phone": "contact_phone",
    "contactphone": "contact_phone",
    "website": "company_website",
    "companywebsite": "company_website",
    "url": "company_website",
    "domain": "company_website",
    "city": "city",
    "state": "state",
    "location": "location",
    "notes": "notes",
    "valueprop": "value_proposition",
    "valueproposition": "value_proposition",
    "description": "value_proposition",
    "linkedin": "linkedin_url",
    "personlinkedin": "linkedin_url",
    "companylinkedin": "company_linkedin_url",
}


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[^a-z0-9-]+", "", text)


def _extract_email(value: Any) -> str:
    match = EMAIL_PATTERN.search(str(value or ""))
    return match.group(0).strip() if match else ""


def _has_phone_candidate(value: Any) -> bool:
    verification = verify_phone_legitimacy(value)
    return str(verification.get("status") or "") not in {"invalid", "no_phone_found"}


def _find_email_column(ws, requested_header: str = "") -> int | None:
    headers = [ws.cell(row=1, column=column).value for column in range(1, ws.max_column + 1)]
    requested = _normalize_header(requested_header)

    if requested:
        for index, header in enumerate(headers, start=1):
            if _normalize_header(header) == requested:
                return index
        return None

    for index, header in enumerate(headers, start=1):
        normalized = _normalize_header(header)
        if normalized in EMAIL_HEADER_ALIASES:
            return index
        if "email" in normalized and "verification" not in normalized and "status" not in normalized:
            return index

    best_column = None
    best_hits = 0
    max_scan_row = min(ws.max_row, 50)
    for column in range(1, ws.max_column + 1):
        hits = 0
        for row in range(1, max_scan_row + 1):
            if _extract_email(ws.cell(row=row, column=column).value):
                hits += 1
        if hits > best_hits:
            best_hits = hits
            best_column = column

    return best_column if best_hits else None


def _find_phone_column(ws, requested_header: str = "") -> int | None:
    headers = [ws.cell(row=1, column=column).value for column in range(1, ws.max_column + 1)]
    requested = _normalize_header(requested_header)

    if requested:
        for index, header in enumerate(headers, start=1):
            if _normalize_header(header) == requested:
                return index
        return None

    for index, header in enumerate(headers, start=1):
        normalized = _normalize_header(header)
        if normalized in PHONE_HEADER_ALIASES:
            return index
        if "phone" in normalized or "mobile" in normalized or normalized in {"tel", "telephone"}:
            return index

    best_column = None
    best_hits = 0
    max_scan_row = min(ws.max_row, 50)
    for column in range(1, ws.max_column + 1):
        hits = 0
        for row in range(2, max_scan_row + 1):
            if _has_phone_candidate(ws.cell(row=row, column=column).value):
                hits += 1
        if hits > best_hits:
            best_hits = hits
            best_column = column

    return best_column if best_hits else None


def _find_website_column(ws) -> int | None:
    headers = [ws.cell(row=1, column=column).value for column in range(1, ws.max_column + 1)]
    for index, header in enumerate(headers, start=1):
        normalized = _normalize_header(header)
        if normalized in WEBSITE_HEADER_ALIASES:
            return index
        if ("website" in normalized or "url" in normalized or "domain" in normalized) and "linkedin" not in normalized:
            return index
    return None


def _ensure_status_column(ws, header: str = DEFAULT_STATUS_HEADER) -> int:
    normalized_target = _normalize_header(header)
    for column in range(1, ws.max_column + 1):
        if _normalize_header(ws.cell(row=1, column=column).value) == normalized_target:
            return column

    column = ws.max_column + 1
    ws.cell(row=1, column=column).value = header
    return column


def _ensure_phone_match_columns(ws) -> dict[str, int]:
    return {key: _ensure_status_column(ws, header) for key, header in PHONE_COMPANY_MATCH_HEADERS.items()}


def _status_header_for_mode(mode: str) -> str:
    return "phone_verification_status" if mode == "phone" else DEFAULT_STATUS_HEADER


def _ensure_summary_column(ws, header: str = DEFAULT_SUMMARY_HEADER) -> int:
    for column in range(1, ws.max_column + 1):
        normalized = _normalize_header(ws.cell(row=1, column=column).value)
        if normalized in SUMMARY_HEADER_ALIASES:
            return column

    column = ws.max_column + 1
    ws.cell(row=1, column=column).value = header
    return column


def _row_to_lead_dict(ws, row: int, source_headers: list[Any]) -> dict[str, Any]:
    lead: dict[str, Any] = {}
    for column, header in enumerate(source_headers, start=1):
        normalized = _normalize_header(header)
        value = ws.cell(row=row, column=column).value
        if value in (None, ""):
            continue

        key = HEADER_TO_LEAD_KEY.get(normalized)
        if not key:
            if "email" in normalized and "verification" not in normalized:
                key = "contact_email"
            elif "phone" in normalized:
                key = "contact_phone"
            elif "website" in normalized or "url" in normalized:
                key = "company_website"
            elif "linkedin" in normalized:
                key = "linkedin_url"
            elif "summary" in normalized:
                key = "lead_summary"
        if key and not lead.get(key):
            lead[key] = value
    return lead


def _write_lead_summary_if_needed(ws, row: int, summary_col: int, source_headers: list[Any]) -> None:
    existing = str(ws.cell(row=row, column=summary_col).value or "").strip()
    if existing:
        return
    lead = _row_to_lead_dict(ws, row, source_headers)
    if not any(lead.get(key) for key in ("company_name", "specialty", "industry", "company_website", "notes")):
        return
    ws.cell(row=row, column=summary_col).value = build_plain_lead_summary(lead)


def _verification_label(verification: dict) -> str:
    status = str(verification.get("status") or "").strip().lower()
    confidence = verification.get("confidence", 0)
    message = str(verification.get("message") or "").strip()

    if status == "valid":
        label = "verified"
    elif status == "risky":
        label = "needs manual review"
    elif status in {"invalid", "error"}:
        label = status
    else:
        label = status or "unknown"

    if confidence not in ("", None):
        label = f"{label} ({confidence}%)"
    if message:
        label = f"{label} - {message}"
    return label


def _phone_verification_label(verification: dict) -> str:
    status = str(verification.get("status") or "").strip().lower()
    confidence = verification.get("confidence", 0)
    message = str(verification.get("message") or "").strip()
    normalized = str(verification.get("normalized") or "").strip()
    e164 = str(verification.get("e164") or "").strip()

    if status == "valid":
        label = "verified phone"
    elif status == "risky":
        label = "needs manual review"
    elif status == "no_phone_found":
        label = "no_phone_found"
    elif status == "invalid":
        label = "invalid phone"
    else:
        label = status or "unknown"

    if confidence not in ("", None):
        label = f"{label} ({confidence}%)"
    if normalized:
        label = f"{label} - {normalized}"
    if e164:
        label = f"{label} / {e164}"
    if message:
        label = f"{label} - {message}"
    return label


def _phone_match_status_label(match: dict | None) -> str:
    if not match:
        return ""
    status = str(match.get("status") or "").strip()
    message = str(match.get("message") or "").strip()
    return f"{status} - {message}" if message else status


def _write_phone_match_result(ws, row: int, match_cols: dict[str, int], match: dict | None) -> None:
    if not match_cols:
        return
    match = match or {}
    scraped_phones = match.get("scraped_phones") or []
    if not isinstance(scraped_phones, list):
        scraped_phones = [scraped_phones]
    values = {
        "status": _phone_match_status_label(match),
        "source": match.get("source_url", ""),
        "confidence": match.get("confidence", ""),
        "matched_phone": match.get("matched_phone", ""),
        "scraped_phones": " | ".join(str(phone).strip() for phone in scraped_phones if str(phone).strip()),
    }
    for key, value in values.items():
        column = match_cols.get(key)
        if column:
            ws.cell(row=row, column=column).value = value


def _load_workbook_from_bytes(data: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is not installed. Run `pip install -r requirements.txt` and restart the server.",
        ) from exc

    try:
        return load_workbook(io.BytesIO(data))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read Excel file: {type(exc).__name__}") from exc


@router.post("/bulk-excel")
async def verify_excel_emails(
    file: UploadFile = File(...),
    email_column: str = Form(""),
    column_header: str = Form(""),
    verification_type: str = Form("email"),
    phone_region: str = Form("US"),
    max_concurrency: int = Form(4),
):
    mode = "phone" if str(verification_type or "").strip().lower() == "phone" else "email"
    filename = file.filename or "emails.xlsx"
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Upload a .xlsx or .xlsm Excel file.")

    data = await file.read()
    if not data:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Excel file is too large. Keep it under 8 MB.")

    workbook = await run_in_threadpool(_load_workbook_from_bytes, data)
    jobs: list[tuple[Any, int, int, int, str, str, dict[str, int]]] = []
    no_value_rows: list[tuple[Any, int, int, dict[str, int]]] = []
    sheet_summaries: list[dict[str, Any]] = []
    requested_header = column_header or email_column
    phone_matcher = PhoneCompanyMatcher() if mode == "phone" else None

    for ws in workbook.worksheets:
        value_col = _find_phone_column(ws, requested_header) if mode == "phone" else _find_email_column(ws, requested_header)
        if not value_col:
            sheet_summaries.append(
                {
                    "sheet": ws.title,
                    "mode": mode,
                    "column_found": False,
                    "email_column_found": False if mode == "email" else None,
                    "phone_column_found": False if mode == "phone" else None,
                    "checks": 0,
                    "emails": 0,
                    "phones": 0,
                }
            )
            continue

        source_headers = [ws.cell(row=1, column=column).value for column in range(1, ws.max_column + 1)]
        website_col = _find_website_column(ws) if mode == "phone" else None
        summary_col = _ensure_summary_column(ws)
        status_col = _ensure_status_column(ws, _status_header_for_mode(mode))
        phone_match_cols = _ensure_phone_match_columns(ws) if mode == "phone" else {}
        count = 0
        for row in range(2, ws.max_row + 1):
            _write_lead_summary_if_needed(ws, row, summary_col, source_headers)
            raw_value = ws.cell(row=row, column=value_col).value
            raw_text = str(raw_value or "").strip()
            if not raw_text or raw_text.lower() in EMPTY_EMAIL_VALUES:
                no_value_rows.append((ws, row, status_col, phone_match_cols))
                continue

            value = raw_text
            website_value = ""
            if website_col:
                website_value = str(ws.cell(row=row, column=website_col).value or "").strip()
            if mode == "email":
                email = _extract_email(raw_text)
                value = email or raw_text
                if not value:
                    no_value_rows.append((ws, row, status_col, phone_match_cols))
                    continue

            jobs.append((ws, row, value_col, status_col, value, website_value, phone_match_cols))
            count += 1
        sheet_summaries.append(
            {
                "sheet": ws.title,
                "mode": mode,
                "column_found": True,
                "email_column_found": True if mode == "email" else None,
                "phone_column_found": True if mode == "phone" else None,
                "checks": count,
                "emails": count if mode == "email" else 0,
                "phones": count if mode == "phone" else 0,
            }
        )

    if not jobs:
        if no_value_rows:
            empty_status = "no_phone_found" if mode == "phone" else "no_email_found"
            for ws, row, status_col, match_cols in no_value_rows:
                ws.cell(row=row, column=status_col).value = empty_status
                if mode == "phone":
                    _write_phone_match_result(
                        ws,
                        row,
                        match_cols,
                        {
                            "status": "no_phone_found",
                            "confidence": "low",
                            "message": "No phone number was available to match against the company website.",
                        },
                    )
            output = io.BytesIO()
            await run_in_threadpool(workbook.save, output)
            output.seek(0)
            download_name = f"verified_{filename.rsplit('.', 1)[0]}.xlsx"
            status_counts = {empty_status: len(no_value_rows)}
            headers = {
                "Content-Disposition": f'attachment; filename="{download_name}"',
                "X-Bulk-Verification-Summary": json.dumps(
                    {
                        "mode": mode,
                        "total_checks": 0,
                        "total_emails": 0,
                        "total_phones": 0,
                        "status_counts": status_counts,
                        "sheets": sheet_summaries,
                    }
                ),
                "X-Email-Verification-Summary": json.dumps(
                    {
                        "mode": mode,
                        "total_checks": 0,
                        "total_emails": 0,
                        "total_phones": 0,
                        "status_counts": status_counts,
                        "sheets": sheet_summaries,
                    }
                ),
            }
            return StreamingResponse(
                output,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers=headers,
            )
        target_label = "phone numbers" if mode == "phone" else "email addresses"
        raise HTTPException(status_code=400, detail=f"No {target_label} were found in the workbook.")

    concurrency = max(1, min(int(max_concurrency or 4), 8))
    semaphore = asyncio.Semaphore(concurrency)

    async def verify_job(job: tuple[Any, int, int, int, str, str, dict[str, int]]):
        async with semaphore:
            ws, row, value_col, status_col, value, website_value, match_cols = job
            phone_match = None
            if mode == "phone":
                verification = await run_in_threadpool(verify_phone_legitimacy, value, region=phone_region or "US")
                if phone_matcher:
                    match_phone = verification.get("e164") or verification.get("normalized") or value
                    phone_match = await phone_matcher.verify(match_phone, website_value, region=phone_region or "US")
            else:
                verification = await run_in_threadpool(verify_email_legitimacy, value)
            return ws, row, value_col, status_col, verification, phone_match, match_cols

    results = await asyncio.gather(*(verify_job(job) for job in jobs))
    status_counts: dict[str, int] = {}

    if no_value_rows:
        empty_status = "no_phone_found" if mode == "phone" else "no_email_found"
        for ws, row, status_col, match_cols in no_value_rows:
            ws.cell(row=row, column=status_col).value = empty_status
            if mode == "phone":
                _write_phone_match_result(
                    ws,
                    row,
                    match_cols,
                    {
                        "status": "no_phone_found",
                        "confidence": "low",
                        "message": "No phone number was available to match against the company website.",
                    },
                )
        status_counts[empty_status] = len(no_value_rows)

    for ws, row, value_col, status_col, verification, phone_match, match_cols in results:
        status = str(verification.get("status") or "unknown").lower()
        summary_status = "needs_manual_review" if status == "risky" else status
        status_counts[summary_status] = status_counts.get(summary_status, 0) + 1
        if mode == "phone" and status in {"valid", "risky"}:
            normalized_phone = str(verification.get("normalized") or verification.get("e164") or "").strip()
            if normalized_phone:
                ws.cell(row=row, column=value_col).value = normalized_phone
        ws.cell(row=row, column=status_col).value = (
            _phone_verification_label(verification) if mode == "phone" else _verification_label(verification)
        )
        if mode == "phone":
            _write_phone_match_result(ws, row, match_cols, phone_match)

    output = io.BytesIO()
    await run_in_threadpool(workbook.save, output)
    output.seek(0)

    summary = {
        "mode": mode,
        "total_checks": len(jobs),
        "total_emails": len(jobs) if mode == "email" else 0,
        "total_phones": len(jobs) if mode == "phone" else 0,
        "status_counts": status_counts,
        "sheets": sheet_summaries,
    }
    download_name = f"verified_{filename.rsplit('.', 1)[0]}.xlsx"
    headers = {
        "Content-Disposition": f'attachment; filename="{download_name}"',
        "X-Bulk-Verification-Summary": json.dumps(summary),
        "X-Email-Verification-Summary": json.dumps(summary),
    }
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )
