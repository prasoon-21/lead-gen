from __future__ import annotations

import asyncio
import contextlib
import json
import os
import re
import time
import uuid
from dataclasses import asdict, dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from core.services.velit.discovery import normalize_domain
from core.services.velit.velit_lead_pipeline import VelitLeadPipeline


DEFAULT_AUTOMATION_ID = "default"
DEFAULT_OUTPUT_FOLDER = "exports/velit_scheduled"
DATA_DIR = Path("data")
AUTOMATION_FILE = DATA_DIR / "velit_batch_automations.json"
QUEUE_FILE = DATA_DIR / "velit_batch_queue.json"
RUNS_FILE = DATA_DIR / "velit_batch_runs.json"
RUN_STATUSES = {"pending", "scheduled", "running", "completed", "failed", "paused", "skipped"}
DEFAULT_AUTOMATION_TIMEOUT_SECONDS = 600


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _iso(value: Optional[datetime] = None) -> str:
    return (value or _utc_now()).isoformat()


def _parse_iso(value: str | None) -> Optional[datetime]:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    except Exception:
        return None


def _slugify(value: str, fallback: str = "item") -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", (value or "").lower()).strip("_")
    return slug or fallback


def _clean_cell(value: Any) -> str:
    if isinstance(value, list):
        return " | ".join(_clean_cell(item) for item in value if _clean_cell(item))
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    text = str(value or "").strip()
    return "" if text.lower() in {"none", "null", "n/a", "na", "-", "--"} else text


@dataclass
class VelitBatchAutomation:
    id: str = DEFAULT_AUTOMATION_ID
    name: str = "Velit multi-location run"
    interval_hours: float = 4.0
    output_folder: str = DEFAULT_OUTPUT_FOLDER
    enabled: bool = False
    running: bool = False
    repeat_mode: bool = False
    run_immediately: bool = True
    current_item_id: Optional[str] = None
    created_at: str = ""
    updated_at: str = ""
    last_run_at: Optional[str] = None
    next_run_at: Optional[str] = None
    completed_at: Optional[str] = None


@dataclass
class VelitQueueItem:
    id: str
    automation_id: str
    order: int
    label: str
    location: str
    target_count: int
    target_role: str
    seed_query: str
    output_folder: str
    status: str
    created_at: str
    updated_at: str
    started_at: Optional[str] = None
    finished_at: Optional[str] = None
    last_error: str = ""
    output_file_path: str = ""
    lead_count: int = 0
    run_count: int = 0
    industry: str = "Van Upfitter"
    duration_seconds: float = 0.0


class VelitBatchScheduler:
    """Persistent in-process queue scheduler for multi-location Velit lead runs."""

    def __init__(self, *, sleep_seconds: int = 30) -> None:
        self.sleep_seconds = max(5, int(sleep_seconds or 30))
        self._lock = asyncio.Lock()
        self._task: Optional[asyncio.Task[Any]] = None
        self._stop_event = asyncio.Event()
        self.automations = self._load_automations()
        self.active_automation_id = DEFAULT_AUTOMATION_ID if DEFAULT_AUTOMATION_ID in self.automations else next(iter(self.automations))
        self.items = self._load_items()
        self.runs = self._load_runs()
        for automation in self.automations.values():
            automation.output_folder = self._ensure_output_folder(automation.output_folder)
        self._repair_loaded_items()

    async def start_runtime(self) -> None:
        if self._task and not self._task.done():
            return
        self._stop_event.clear()
        self._task = asyncio.create_task(self._loop())

    async def stop_runtime(self) -> None:
        self._stop_event.set()
        if self._task:
            self._task.cancel()
            try:
                await self._task
            except asyncio.CancelledError:
                pass
        self._task = None

    async def _loop(self) -> None:
        while not self._stop_event.is_set():
            try:
                await self.run_due_item()
            except asyncio.CancelledError:
                raise
            except Exception:
                pass
            try:
                await asyncio.wait_for(self._stop_event.wait(), timeout=self.sleep_seconds)
            except asyncio.TimeoutError:
                pass

    async def get_snapshot(self, automation_id: Optional[str] = None) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            return self._snapshot_unlocked(automation.id)

    async def create_automation(
        self,
        *,
        name: str = "",
        interval_hours: Optional[float] = None,
        output_folder: str = "",
        repeat_mode: bool = False,
        run_immediately: bool = True,
    ) -> Dict[str, Any]:
        now = _iso()
        automation_id = uuid.uuid4().hex
        automation = VelitBatchAutomation(
            id=automation_id,
            name=str(name or "").strip() or "Velit multi-location run",
            interval_hours=self._validate_interval(interval_hours if interval_hours is not None else 4.0),
            output_folder=self._ensure_output_folder(output_folder or DEFAULT_OUTPUT_FOLDER),
            repeat_mode=bool(repeat_mode),
            run_immediately=bool(run_immediately),
            created_at=now,
            updated_at=now,
        )
        async with self._lock:
            self.automations[automation.id] = automation
            self.active_automation_id = automation.id
            self._save_all()
            return self._snapshot_unlocked(automation.id)

    async def delete_automation(self, automation_id: str) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            if automation.running:
                raise ValueError("Cannot delete an automation while it is running.")
            if len(self.automations) <= 1:
                raise ValueError("Cannot delete the last automation.")
            item_ids = {item.id for item in self.items if item.automation_id == automation.id}
            self.items = [item for item in self.items if item.automation_id != automation.id]
            self.runs = [
                run
                for run in self.runs
                if run.get("automation_id") != automation.id and run.get("item_id") not in item_ids
            ]
            del self.automations[automation.id]
            self.active_automation_id = DEFAULT_AUTOMATION_ID if DEFAULT_AUTOMATION_ID in self.automations else next(iter(self.automations))
            self._save_all()
            return self._snapshot_unlocked(self.active_automation_id)

    async def update_automation(
        self,
        *,
        automation_id: Optional[str] = None,
        name: Optional[str] = None,
        interval_hours: Optional[float] = None,
        output_folder: Optional[str] = None,
        repeat_mode: Optional[bool] = None,
        run_immediately: Optional[bool] = None,
    ) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            if interval_hours is not None:
                automation.interval_hours = self._validate_interval(interval_hours)
            if name is not None:
                automation.name = str(name or "").strip() or "Velit multi-location run"
            if output_folder is not None:
                automation.output_folder = self._ensure_output_folder(output_folder or DEFAULT_OUTPUT_FOLDER)
            if repeat_mode is not None:
                automation.repeat_mode = bool(repeat_mode)
            if run_immediately is not None:
                automation.run_immediately = bool(run_immediately)
            automation.updated_at = _iso()
            self._save_all()
            return self._snapshot_unlocked(automation.id)

    async def start_automation(self, automation_id: Optional[str] = None) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            automation.enabled = True
            automation.completed_at = None
            automation.updated_at = _iso()
            if automation.run_immediately:
                automation.next_run_at = _iso()
            elif not automation.next_run_at or _parse_iso(automation.next_run_at) <= _utc_now():
                automation.next_run_at = _iso(_utc_now() + self._interval_delta(automation))
            self._save_all()
            snapshot = self._snapshot_unlocked(automation.id)
            should_run = automation.run_immediately
            selected_id = automation.id
        if should_run:
            asyncio.create_task(self.run_due_item(selected_id))
        return snapshot

    async def pause_automation(self, automation_id: Optional[str] = None) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            automation.enabled = False
            automation.updated_at = _iso()
            self._save_all()
            return self._snapshot_unlocked(automation.id)

    async def resume_automation(self, automation_id: Optional[str] = None) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            automation.enabled = True
            automation.completed_at = None
            automation.updated_at = _iso()
            if not automation.next_run_at:
                automation.next_run_at = _iso(_utc_now() + self._interval_delta(automation))
            self._save_all()
            return self._snapshot_unlocked(automation.id)

    async def stop_automation(self, automation_id: Optional[str] = None) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            automation.enabled = False
            if not automation.running:
                automation.current_item_id = None
            automation.updated_at = _iso()
            self._save_all()
            return self._snapshot_unlocked(automation.id)

    async def reset_timer(self, automation_id: Optional[str] = None) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            if automation.running:
                raise ValueError("Cannot reset timer while automation is running.")
            if automation.enabled:
                automation.next_run_at = _iso(_utc_now() + self._interval_delta(automation))
                automation.completed_at = None
                automation.updated_at = _iso()
                self._save_all()
            return self._snapshot_unlocked(automation.id)

    async def add_item(
        self,
        *,
        automation_id: Optional[str] = None,
        label: str = "",
        location: str,
        target_count: int,
        target_role: str = "Founder",
        industry: str = "Van Upfitter",
        seed_query: str = "",
        output_folder: str = "",
    ) -> Dict[str, Any]:
        location = self._validate_location(location)
        target_count = self._validate_target_count(target_count)
        target_role = self._validate_role(target_role)
        industry = self._validate_industry(industry)
        now = _iso()
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            folder = self._ensure_output_folder(output_folder or automation.output_folder or DEFAULT_OUTPUT_FOLDER)
            item = VelitQueueItem(
                id=uuid.uuid4().hex,
                automation_id=automation.id,
                order=self._next_order_unlocked(automation.id),
                label=str(label or "").strip() or f"{location} {target_role} leads",
                location=location,
                target_count=target_count,
                target_role=target_role,
                industry=industry,
                seed_query=str(seed_query or "").strip(),
                output_folder=folder,
                status="pending",
                created_at=now,
                updated_at=now,
            )
            self.items.append(item)
            self._sort_items_unlocked(automation.id)
            self._save_all()
            return asdict(item)

    async def list_items(self, automation_id: Optional[str] = None) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            return {"items": [asdict(item) for item in self._sorted_items(automation.id)]}

    async def update_item(self, item_id: str, **updates: Any) -> Dict[str, Any]:
        async with self._lock:
            item = self._get_item_unlocked(item_id)
            if item.status == "running":
                raise ValueError("Cannot update an item while it is running.")
            if "label" in updates and updates["label"] is not None:
                item.label = str(updates["label"] or "").strip() or item.label
            if "location" in updates and updates["location"] is not None:
                item.location = self._validate_location(updates["location"])
            if "target_count" in updates and updates["target_count"] is not None:
                item.target_count = self._validate_target_count(updates["target_count"])
            if "target_role" in updates and updates["target_role"] is not None:
                item.target_role = self._validate_role(updates["target_role"])
            if "industry" in updates and updates["industry"] is not None:
                item.industry = self._validate_industry(updates["industry"])
            if "seed_query" in updates and updates["seed_query"] is not None:
                item.seed_query = str(updates["seed_query"] or "").strip()
            if "output_folder" in updates and updates["output_folder"] is not None:
                automation = self._automation_unlocked(item.automation_id)
                item.output_folder = self._ensure_output_folder(updates["output_folder"] or automation.output_folder)
            item.updated_at = _iso()
            self._save_all()
            return asdict(item)

    async def delete_item(self, item_id: str) -> Dict[str, Any]:
        async with self._lock:
            item = self._get_item_unlocked(item_id)
            if item.status == "running":
                raise ValueError("Cannot delete an item while it is running.")
            self.items = [candidate for candidate in self.items if candidate.id != item_id]
            self._normalize_orders_unlocked(item.automation_id)
            self._save_all()
            return self._snapshot_unlocked(item.automation_id)

    async def reset_item(self, item_id: str) -> Dict[str, Any]:
        async with self._lock:
            item = self._get_item_unlocked(item_id)
            if item.status == "running":
                raise ValueError("Cannot reset an item while it is running.")
            item.status = "pending"
            item.started_at = None
            item.finished_at = None
            item.last_error = ""
            item.output_file_path = ""
            item.lead_count = 0
            item.duration_seconds = 0.0
            item.updated_at = _iso()
            self._save_all()
            return asdict(item)

    async def reorder_items(self, item_ids: List[str], automation_id: Optional[str] = None) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            known = {item.id: item for item in self.items if item.automation_id == automation.id}
            requested = [item_id for item_id in item_ids if item_id in known]
            missing = [item for item in self._sorted_items(automation.id) if item.id not in requested]
            ordered = [known[item_id] for item_id in requested] + missing
            for index, item in enumerate(ordered, start=1):
                item.order = index
                item.updated_at = _iso()
            other_items = [item for item in self.items if item.automation_id != automation.id]
            self.items = other_items + ordered
            self._save_all()
            return self._snapshot_unlocked(automation.id)

    async def run_item_now(self, item_id: str) -> Dict[str, Any]:
        async with self._lock:
            item = self._get_item_unlocked(item_id)
            if self._any_running_unlocked():
                raise ValueError("Another queue item is already running.")
        asyncio.create_task(self._run_item(item_id, manual=True))
        return await self.get_snapshot(item.automation_id)

    async def run_due_item(self, automation_id: Optional[str] = None) -> Dict[str, Any]:
        item_id = ""
        async with self._lock:
            if automation_id and automation_id not in self.automations:
                return self._snapshot_unlocked()
            automation_ids = [self._automation_unlocked(automation_id).id] if automation_id else list(self.automations.keys())
            selected_automation = self._automation_unlocked(automation_ids[0])
            if self._any_running_unlocked():
                return self._snapshot_unlocked(selected_automation.id)
            selected_item: Optional[VelitQueueItem] = None
            selected_automation_id = selected_automation.id
            for candidate_id in automation_ids:
                automation = self._automation_unlocked(candidate_id)
                if not automation.enabled or automation.running:
                    continue
                next_run_at = _parse_iso(automation.next_run_at)
                if next_run_at and next_run_at > _utc_now():
                    continue
                item = self._next_pending_item_unlocked(automation.id)
                if not item:
                    automation_items = [candidate for candidate in self.items if candidate.automation_id == automation.id]
                    if automation.repeat_mode and automation_items:
                        for candidate in automation_items:
                            if candidate.status == "completed":
                                candidate.status = "pending"
                                candidate.updated_at = _iso()
                        item = self._next_pending_item_unlocked(automation.id)
                    if not item:
                        automation.enabled = False
                        automation.completed_at = _iso()
                        automation.updated_at = _iso()
                        self._save_all()
                        continue
                selected_item = item
                selected_automation_id = automation.id
                break
            if not selected_item:
                return self._snapshot_unlocked(selected_automation_id)
            item_id = selected_item.id
        await self._run_item(item_id, manual=False)
        return await self.get_snapshot(selected_automation_id)

    async def recent_runs(self, limit: int = 50, automation_id: Optional[str] = None) -> Dict[str, Any]:
        async with self._lock:
            automation = self._automation_unlocked(automation_id)
            runs = [run for run in self.runs if run.get("automation_id") == automation.id]
            return {"runs": list(reversed(runs[-max(1, min(int(limit or 50), 200)) :]))}

    async def _run_item(self, item_id: str, *, manual: bool) -> None:
        async with self._lock:
            if self._any_running_unlocked():
                return
            item = self._get_item_unlocked(item_id)
            automation = self._automation_unlocked(item.automation_id)
            if item.status == "running":
                return
            started_at = _iso()
            automation.running = True
            automation.current_item_id = item.id
            automation.updated_at = started_at
            item.status = "running"
            item.started_at = started_at
            item.finished_at = None
            item.last_error = ""
            item.updated_at = started_at
            self._save_all()

        lead_count = 0
        output_path = ""
        error = ""
        result: Dict[str, Any] = {}
        timed_out = False
        duration_seconds = 0.0
        run_start_time = time.monotonic()
        try:
            pipeline = VelitLeadPipeline()
            timeout_seconds = self._automation_timeout_seconds()
            pipeline_target_count = item.target_count
            run_task = asyncio.create_task(
                pipeline.run(
                    industry=item.industry or "Van Upfitter",
                    location=item.location,
                    seed_query=item.seed_query,
                    target_count=pipeline_target_count,
                )
            )
            done, _ = await asyncio.wait({run_task}, timeout=timeout_seconds)
            timed_out = run_task not in done
            if timed_out:
                timeout_message = f"TimeoutError: Automation exceeded {timeout_seconds} seconds and saved partial leads."
                result = pipeline.get_partial_result(error=timeout_message, status="timeout")
                run_task.cancel()
                with contextlib.suppress(asyncio.CancelledError):
                    await run_task
            else:
                result = run_task.result()

            metadata = result.setdefault("metadata", {})
            metadata.update(
                {
                    "automation_search_mode": "normal_with_timeout",
                    "automation_timed_out": timed_out,
                    "automation_timeout_seconds": timeout_seconds,
                    "automation_requested_target_count": item.target_count,
                    "automation_pipeline_target_count": pipeline_target_count,
                }
            )
            raw_leads = list(result.get("leads") or [])
            leads, filter_stats = self._filter_export_leads(raw_leads, target_role=item.target_role)
            result["leads"] = leads
            metadata["automation_export_filter"] = filter_stats
            if result.get("error"):
                error = str(result.get("error"))
            lead_count = len(leads)
            if timed_out and lead_count > 0:
                metadata["automation_completed_with_partial_timeout"] = True
                error = ""
            duration_seconds = round(time.monotonic() - run_start_time, 2)
            metadata["duration_seconds"] = duration_seconds
            output_path = self._write_xlsx(item=item, leads=leads, result=result, started_at=started_at)
        except asyncio.TimeoutError:
            timeout_seconds = self._automation_timeout_seconds()
            error = f"TimeoutError: Automation exceeded {timeout_seconds} seconds and was stopped."
            duration_seconds = round(time.monotonic() - run_start_time, 2)
            result = {
                "leads": [],
                "steps": [],
                "error": error,
                "metadata": {
                    "pipeline": "velit_specialized_pipeline",
                    "status": "timeout",
                    "automation_search_mode": "normal_with_timeout",
                    "automation_timeout_seconds": timeout_seconds,
                    "automation_requested_target_count": item.target_count,
                    "automation_pipeline_target_count": item.target_count,
                    "duration_seconds": duration_seconds,
                    "automation_export_filter": {
                        "input_count": 0,
                        "exported_count": 0,
                        "dropped_missing_company_name": 0,
                        "dropped_duplicate_company": 0,
                    },
                },
            }
            output_path = self._write_xlsx(item=item, leads=[], result=result, started_at=started_at)
        except Exception as exc:
            error = f"{type(exc).__name__}: {exc}"
            duration_seconds = round(time.monotonic() - run_start_time, 2)
            result = {
                "leads": [],
                "steps": [],
                "error": error,
                "metadata": {
                    "pipeline": "velit_specialized_pipeline",
                    "status": "failed",
                    "automation_search_mode": "normal_with_timeout",
                    "automation_requested_target_count": item.target_count,
                    "automation_pipeline_target_count": item.target_count,
                    "duration_seconds": duration_seconds,
                },
            }
            try:
                output_path = self._write_xlsx(item=item, leads=[], result=result, started_at=started_at)
            except Exception:
                output_path = ""

        finished_at = _iso()
        if not duration_seconds:
            duration_seconds = round(time.monotonic() - run_start_time, 2)
        async with self._lock:
            item = self._get_item_unlocked(item_id)
            automation = self._automation_unlocked(item.automation_id)
            item.finished_at = finished_at
            item.updated_at = finished_at
            item.run_count += 1
            item.lead_count = lead_count
            item.output_file_path = output_path
            item.last_error = error
            item.duration_seconds = duration_seconds
            item.status = "failed" if error else "completed"

            automation.running = False
            automation.current_item_id = None
            automation.last_run_at = finished_at
            if automation.enabled:
                automation.next_run_at = _iso(_utc_now() + self._interval_delta(automation))
            automation.updated_at = finished_at

            self.runs.append(
                {
                    "id": uuid.uuid4().hex,
                    "automation_id": automation.id,
                    "automation_name": automation.name,
                    "item_id": item.id,
                    "order": item.order,
                    "label": item.label,
                    "location": item.location,
                    "industry": item.industry,
                    "target_role": item.target_role,
                    "target_count": item.target_count,
                    "status": item.status,
                    "manual": manual,
                    "started_at": started_at,
                    "finished_at": finished_at,
                    "duration_seconds": duration_seconds,
                    "lead_count": lead_count,
                    "output_file_path": output_path,
                    "error": error,
                    "metadata": result.get("metadata", {}) if isinstance(result, dict) else {},
                }
            )
            self.runs = self.runs[-500:]
            self._save_all()

    def _write_xlsx(
        self,
        *,
        item: VelitQueueItem,
        leads: List[Dict[str, Any]],
        result: Dict[str, Any],
        started_at: str,
    ) -> str:
        automation = self._automation_unlocked(item.automation_id)
        output_folder = Path(self._ensure_output_folder(item.output_folder or automation.output_folder))
        timestamp = _utc_now().strftime("%Y%m%d_%H%M%S")
        filename = f"{item.order:03d}_{_slugify(item.location, 'location')}_{_slugify(item.target_role, 'role')}_{timestamp}.xlsx"
        path = output_folder / filename

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Leads"
        headers = [
            "Company Name",
            "Contact Name",
            "Target Role",
            "Contact Title",
            "Email",
            "Phone",
            "Website",
            "Person LinkedIn",
            "Company LinkedIn",
            "Location",
            "City",
            "State",
            "Lead Summary",
            "Value Proposition",
            "Quality Score",
            "Verification Status",
            "Phone Status",
            "Source",
            "Generated At",
        ]
        sheet.append(headers)
        generated_at = _iso()
        for lead in leads:
            sheet.append(
                [
                    _clean_cell(lead.get("company_name")),
                    _clean_cell(lead.get("contact_person_name") or lead.get("founder_name")),
                    item.target_role,
                    _clean_cell(lead.get("contact_person_title")),
                    _clean_cell(lead.get("contact_email")),
                    _clean_cell(lead.get("contact_phone")),
                    _clean_cell(lead.get("company_website")),
                    _clean_cell(lead.get("linkedin_url")),
                    _clean_cell(lead.get("company_linkedin_url")),
                    _clean_cell(lead.get("location") or item.location),
                    _clean_cell(lead.get("city")),
                    _clean_cell(lead.get("state")),
                    _clean_cell(lead.get("lead_summary")),
                    _clean_cell(lead.get("value_proposition")),
                    _clean_cell(lead.get("quality_score")),
                    _clean_cell(lead.get("verification_status")),
                    _clean_cell(lead.get("phone_validation_status")),
                    _clean_cell(lead.get("source")),
                    generated_at,
                ]
            )

        metadata = workbook.create_sheet("Run Metadata")
        finished_at = _iso()
        duration_seconds = (result.get("metadata") or {}).get("duration_seconds", "")
        duration_human = self._format_duration(float(duration_seconds)) if duration_seconds else ""
        metadata_rows = [
            ("Automation ID", automation.id),
            ("Automation Name", automation.name),
            ("Queue Item ID", item.id),
            ("Queue Order", item.order),
            ("Label", item.label),
            ("Location", item.location),
            ("Industry", item.industry),
            ("Target count", item.target_count),
            ("Target role", item.target_role),
            ("Seed query", item.seed_query),
            ("Started at", started_at),
            ("Finished at", finished_at),
            ("Duration (seconds)", duration_seconds),
            ("Duration", duration_human),
            ("Lead count", len(leads)),
            ("Output file path", str(path)),
            ("Pipeline name", (result.get("metadata") or {}).get("pipeline", "velit_specialized_pipeline")),
            ("Automation search mode", (result.get("metadata") or {}).get("automation_search_mode", "")),
            ("Automation max runtime seconds", (result.get("metadata") or {}).get("automation_timeout_seconds", "")),
            ("Requested target count", (result.get("metadata") or {}).get("automation_requested_target_count", "")),
            ("Pipeline target count", (result.get("metadata") or {}).get("automation_pipeline_target_count", "")),
            ("Pipeline error", result.get("error", "")),
        ]
        export_filter = (result.get("metadata") or {}).get("automation_export_filter")
        if isinstance(export_filter, dict):
            metadata_rows.extend(
                [
                    ("Raw leads before export filter", export_filter.get("input_count", "")),
                    ("Rows exported after filter", export_filter.get("exported_count", "")),
                    ("Rows dropped missing company name", export_filter.get("dropped_missing_company_name", "")),
                    ("Rows dropped duplicate company", export_filter.get("dropped_duplicate_company", "")),
                ]
            )
        for row in metadata_rows:
            metadata.append(row)

        self._style_sheet(sheet)
        self._style_sheet(metadata)
        workbook.save(path)
        return str(path)

    @staticmethod
    def _style_sheet(sheet: Any) -> None:
        header_fill = PatternFill("solid", fgColor="D9EAF7")
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            width = max(12, min(max_length + 2, 48))
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        sheet.freeze_panes = "A2"
        if sheet.max_column and sheet.max_row:
            sheet.auto_filter.ref = sheet.dimensions

    @staticmethod
    def _format_duration(seconds: float) -> str:
        total = int(seconds)
        def unit(value: int, singular: str) -> str:
            return f"{value} {singular}{'' if value == 1 else 's'}"
        if total < 60:
            return unit(total, "second")
        minutes, secs = divmod(total, 60)
        if minutes < 60:
            return f"{unit(minutes, 'minute')} {unit(secs, 'second')}"
        hours, minutes = divmod(minutes, 60)
        return f"{unit(hours, 'hour')} {unit(minutes, 'minute')} {unit(secs, 'second')}"

    @staticmethod
    def _automation_timeout_seconds() -> int:
        return VelitBatchScheduler._bounded_int_env(
            "VELIT_AUTOMATION_MAX_RUNTIME_SECONDS",
            default=DEFAULT_AUTOMATION_TIMEOUT_SECONDS,
            minimum=60,
            maximum=600,
        )

    @classmethod
    def _filter_export_leads(cls, leads: List[Dict[str, Any]], *, target_role: str) -> tuple[List[Dict[str, Any]], Dict[str, int]]:
        filtered: List[Dict[str, Any]] = []
        seen_keys: set[str] = set()
        stats = {
            "input_count": len(leads or []),
            "exported_count": 0,
            "dropped_missing_company_name": 0,
            "dropped_duplicate_company": 0,
        }
        for lead in leads or []:
            if not isinstance(lead, dict):
                stats["dropped_missing_company_name"] += 1
                continue
            item = dict(lead)
            company_name = cls._resolve_export_company_name(item)
            if not cls._is_usable_company_name(company_name):
                stats["dropped_missing_company_name"] += 1
                continue
            item["company_name"] = company_name
            person_name = _clean_cell(item.get("contact_person_name") or item.get("founder_name"))
            if not person_name:
                person_name = cls._person_name_from_email(item.get("contact_email"))
            if person_name:
                item["contact_person_name"] = person_name
                item["founder_name"] = item.get("founder_name") or person_name
            item["contact_person_title"] = _clean_cell(item.get("contact_person_title")) or target_role
            company_key = cls._company_dedupe_key(item)
            if company_key in seen_keys:
                stats["dropped_duplicate_company"] += 1
                continue
            seen_keys.add(company_key)
            filtered.append(item)
        stats["exported_count"] = len(filtered)
        return filtered, stats

    @classmethod
    def _resolve_export_company_name(cls, lead: Dict[str, Any]) -> str:
        website = _clean_cell(lead.get("company_website") or lead.get("website_url") or lead.get("contact_page"))
        email = _clean_cell(lead.get("contact_email") or lead.get("email"))
        title = _clean_cell(lead.get("title") or lead.get("page_title") or lead.get("source_title"))
        source_details = lead.get("source_details") if isinstance(lead.get("source_details"), list) else []
        for detail in source_details:
            if not isinstance(detail, dict):
                continue
            title = title or _clean_cell(detail.get("title"))
            website = website or _clean_cell(detail.get("url"))
            if title and website:
                break
        current = _clean_cell(lead.get("company_name") or lead.get("company"))
        try:
            resolved = VelitLeadPipeline._resolve_company_name(current, title=title, url=website, email=email)
        except Exception:
            resolved = current
        if cls._is_usable_company_name(resolved):
            return resolved
        for candidate in (
            current,
            VelitLeadPipeline._company_name_from_email(email),
            VelitLeadPipeline._company_name_from_url(website),
        ):
            if cls._is_usable_company_name(candidate):
                return candidate
        return ""

    @staticmethod
    def _is_usable_company_name(value: Any) -> bool:
        cleaned = _clean_cell(value)
        lowered = re.sub(r"[^a-z0-9]+", " ", cleaned.lower()).strip()
        if not cleaned or lowered in {"unknown", "unknown company", "n a", "na", "none", "null"}:
            return False
        if VelitLeadPipeline._is_disallowed_company_name(cleaned):
            return False
        return VelitLeadPipeline._looks_like_valid_company_name(cleaned)

    @staticmethod
    def _person_name_from_email(value: Any) -> str:
        email = _clean_cell(value).lower()
        if "@" not in email:
            return ""
        local = email.split("@", 1)[0]
        if local in {"info", "support", "contact", "hello", "sales", "admin", "office", "team", "help", "service", "parts"}:
            return ""
        local = re.sub(r"\+.*$", "", local)
        parts = [part for part in re.split(r"[._-]+", local) if part and not part.isdigit()]
        if len(parts) < 2 or any(len(part) < 2 for part in parts[:2]):
            return ""
        return " ".join(part.capitalize() for part in parts[:3])

    @staticmethod
    def _company_dedupe_key(lead: Dict[str, Any]) -> str:
        website = _clean_cell(lead.get("company_website") or lead.get("contact_page"))
        try:
            domain = normalize_domain(website)
        except Exception:
            domain = ""
        if domain:
            return f"domain:{domain}"
        return "name:" + re.sub(r"[^a-z0-9]+", "", _clean_cell(lead.get("company_name")).lower())

    def _load_automations(self) -> Dict[str, VelitBatchAutomation]:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        now = _iso()
        fallback = VelitBatchAutomation(created_at=now, updated_at=now)
        if AUTOMATION_FILE.exists():
            try:
                data = json.loads(AUTOMATION_FILE.read_text(encoding="utf-8"))
                automations: Dict[str, VelitBatchAutomation] = {}
                raw_records: List[Dict[str, Any]] = []
                if isinstance(data, list):
                    raw_records = [item for item in data if isinstance(item, dict)]
                elif isinstance(data, dict) and any(key in data for key in ("id", "name", "interval_hours")):
                    raw_records = [data]
                elif isinstance(data, dict):
                    raw_records = [item for item in data.values() if isinstance(item, dict)]
                for raw in raw_records:
                    defaults = asdict(VelitBatchAutomation(created_at=now, updated_at=now))
                    for key, value in raw.items():
                        if key in defaults and value is not None:
                            defaults[key] = value
                    automation = VelitBatchAutomation(**defaults)
                    automation.id = str(automation.id or uuid.uuid4().hex).strip() or uuid.uuid4().hex
                    automations[automation.id] = automation
                if automations:
                    return automations
            except Exception:
                pass
        return {fallback.id: fallback}

    def _load_items(self) -> List[VelitQueueItem]:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        if not QUEUE_FILE.exists():
            return []
        try:
            raw_items = json.loads(QUEUE_FILE.read_text(encoding="utf-8"))
        except Exception:
            return []
        items: List[VelitQueueItem] = []
        if isinstance(raw_items, list):
            for raw in raw_items:
                if not isinstance(raw, dict):
                    continue
                try:
                    values = {key: raw.get(key) for key in VelitQueueItem.__dataclass_fields__}
                    values["automation_id"] = values.get("automation_id") or DEFAULT_AUTOMATION_ID
                    values["industry"] = self._validate_industry(values.get("industry") or "Van Upfitter")
                    values["duration_seconds"] = float(values.get("duration_seconds") or 0.0)
                    items.append(VelitQueueItem(**values))
                except Exception:
                    continue
        items.sort(key=lambda item: item.order)
        return items

    def _load_runs(self) -> List[Dict[str, Any]]:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        if not RUNS_FILE.exists():
            return []
        try:
            runs = json.loads(RUNS_FILE.read_text(encoding="utf-8"))
            return runs if isinstance(runs, list) else []
        except Exception:
            return []

    def _save_all(self) -> None:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        AUTOMATION_FILE.write_text(
            json.dumps({automation_id: asdict(automation) for automation_id, automation in self.automations.items()}, indent=2),
            encoding="utf-8",
        )
        QUEUE_FILE.write_text(json.dumps([asdict(item) for item in self._sorted_items()], indent=2), encoding="utf-8")
        RUNS_FILE.write_text(json.dumps(self.runs, indent=2), encoding="utf-8")

    def _snapshot_unlocked(self, automation_id: Optional[str] = None) -> Dict[str, Any]:
        automation = self._automation_unlocked(automation_id)
        items = self._sorted_items(automation.id)
        runs = [run for run in self.runs if run.get("automation_id") == automation.id]
        status_counts: Dict[str, int] = {status: 0 for status in sorted(RUN_STATUSES)}
        for item in items:
            status_counts[item.status] = status_counts.get(item.status, 0) + 1
        return {
            "automation": asdict(automation),
            "automations": [asdict(item) for item in self._sorted_automations()],
            "items": [asdict(item) for item in items],
            "summary": {
                "total_items": len(items),
                "status_counts": status_counts,
                "recent_run_count": len(runs),
            },
            "recent_runs": list(reversed(runs[-25:])),
        }

    def _automation_unlocked(self, automation_id: Optional[str] = None) -> VelitBatchAutomation:
        selected_id = str(automation_id or self.active_automation_id or DEFAULT_AUTOMATION_ID).strip()
        automation = self.automations.get(selected_id)
        if not automation:
            raise KeyError(f"Automation not found: {selected_id}")
        self.active_automation_id = automation.id
        return automation

    def _sorted_automations(self) -> List[VelitBatchAutomation]:
        return sorted(self.automations.values(), key=lambda item: (item.created_at or "", item.name.lower()))

    def _repair_loaded_items(self) -> None:
        changed = False
        for item in self.items:
            if item.automation_id not in self.automations:
                item.automation_id = DEFAULT_AUTOMATION_ID if DEFAULT_AUTOMATION_ID in self.automations else self.active_automation_id
                changed = True
        for automation_id in self.automations:
            self._normalize_orders_unlocked(automation_id)
        if changed:
            self._save_all()

    def _any_running_unlocked(self) -> bool:
        return any(automation.running for automation in self.automations.values())

    def _next_order_unlocked(self, automation_id: Optional[str] = None) -> int:
        automation = self._automation_unlocked(automation_id)
        return max([item.order for item in self.items if item.automation_id == automation.id], default=0) + 1

    def _sort_items_unlocked(self, automation_id: Optional[str] = None) -> None:
        if automation_id is None:
            self.items.sort(key=lambda item: (item.automation_id, item.order))
            return
        automation = self._automation_unlocked(automation_id)
        self.items.sort(key=lambda item: (0 if item.automation_id == automation.id else 1, item.automation_id, item.order))

    def _sorted_items(self, automation_id: Optional[str] = None) -> List[VelitQueueItem]:
        if automation_id is None:
            return sorted(self.items, key=lambda item: (item.automation_id, item.order))
        automation = self._automation_unlocked(automation_id)
        return sorted([item for item in self.items if item.automation_id == automation.id], key=lambda item: item.order)

    def _normalize_orders_unlocked(self, automation_id: Optional[str] = None) -> None:
        items = self._sorted_items(automation_id)
        for index, item in enumerate(items, start=1):
            item.order = index
            item.updated_at = _iso()

    def _get_item_unlocked(self, item_id: str) -> VelitQueueItem:
        for item in self.items:
            if item.id == item_id:
                return item
        raise KeyError(f"Queue item not found: {item_id}")

    def _next_pending_item_unlocked(self, automation_id: Optional[str] = None) -> Optional[VelitQueueItem]:
        for item in self._sorted_items(automation_id):
            if item.status in {"pending", "scheduled"}:
                return item
        return None

    def _interval_delta(self, automation: VelitBatchAutomation) -> timedelta:
        return timedelta(hours=automation.interval_hours)

    @staticmethod
    def _validate_location(value: str) -> str:
        cleaned = " ".join((value or "").split()).strip()
        if len(cleaned) < 2 or len(cleaned) > 120:
            raise ValueError("Location must be 2-120 characters.")
        return cleaned

    @staticmethod
    def _validate_role(value: str) -> str:
        cleaned = " ".join((value or "Founder").split()).strip() or "Founder"
        if len(cleaned) > 120:
            raise ValueError("Target role must be 120 characters or fewer.")
        return cleaned

    @staticmethod
    def _validate_industry(value: str) -> str:
        cleaned = " ".join((value or "Van Upfitter").split()).strip() or "Van Upfitter"
        if len(cleaned) > 120:
            raise ValueError("Industry must be 120 characters or fewer.")
        return cleaned

    @staticmethod
    def _validate_target_count(value: int) -> int:
        try:
            numeric = int(value)
        except Exception as exc:
            raise ValueError("Target count must be a number.") from exc
        if numeric < 1 or numeric > 100:
            raise ValueError("Target count must be between 1 and 100.")
        return numeric

    @staticmethod
    def _validate_interval(value: float) -> float:
        try:
            numeric = float(value)
        except Exception as exc:
            raise ValueError("Interval hours must be a number.") from exc
        if numeric < 0.25 or numeric > 168:
            raise ValueError("Interval hours must be between 0.25 and 168.")
        return numeric

    @staticmethod
    def _ensure_output_folder(value: str) -> str:
        folder = Path(value or DEFAULT_OUTPUT_FOLDER).expanduser()
        if folder.exists() and not folder.is_dir():
            raise ValueError("Output folder points to a file.")
        folder.mkdir(parents=True, exist_ok=True)
        return str(folder)

    @staticmethod
    def _bounded_int_env(name: str, *, default: int, minimum: int, maximum: int) -> int:
        try:
            value = int(os.getenv(name, str(default)))
        except ValueError:
            value = default
        return max(minimum, min(value, maximum))
