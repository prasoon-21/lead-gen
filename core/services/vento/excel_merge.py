from __future__ import annotations

import csv
import io
import json
from copy import copy
from pathlib import Path
from typing import Any, Dict, Iterable, List, Set, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.services.vento.discovery import normalize_social_handle
from core.services.vento.qualification import creator_identity_keys


_HEADER_TO_FIELD = {
    "creator id": "creator_id",
    "batch id": "import_batch_id",
    "creator name": "creator_name",
    "name": "creator_name",
    "email": "email",
    "email verification": "email_verification_status",
    "email source": "email_source_url",
    "email source type": "email_source_type",
    "instagram handle": "instagram_handle",
    "instagram": "instagram_handle",
    "instagram url": "instagram_url",
    "instagram status": "instagram_profile_status",
    "tiktok handle": "tiktok_handle",
    "tiktok": "tiktok_handle",
    "tiktok url": "tiktok_url",
    "tiktok status": "tiktok_profile_status",
    "youtube channel": "youtube_channel",
    "youtube": "youtube_channel",
    "youtube url": "youtube_url",
    "youtube status": "youtube_profile_status",
    "phone": "phone",
    "phone type": "phone_type",
    "website": "website",
    "contact page": "contact_page_url",
    "media kit": "media_kit_url",
    "management contact": "management_contact",
    "location": "location",
    "city": "location_city",
    "state": "location_state",
    "country": "location_country",
    "region": "location_region",
    "location confidence": "location_confidence",
    "location evidence": "location_evidence",
    "location source": "location_source_url",
    "creator type": "creator_type",
    "lead category": "lead_category",
    "identity confidence": "identity_confidence",
    "product fit level": "product_fit_level",
    "product fit score": "product_fit_score",
    "product fit tags": "product_fit_tags",
    "product fit evidence": "product_fit_evidence",
    "niche": "niche",
    "niche tags": "niche_tags",
    "niche verification": "niche_verification_status",
    "follower count": "follower_count",
    "follower status": "follower_count_status",
    "follower count status": "follower_count_status",
    "follower tier": "estimated_follower_tier",
    "follower source": "follower_source_url",
    "engagement rate": "engagement_rate",
    "average views": "average_views",
    "last post date": "last_post_date",
    "quality level": "quality_level",
    "quality reasons": "quality_reasons",
    "review reasons": "review_reasons",
    "relevance score": "relevance_score",
    "contactability score": "contactability_score",
    "verification score": "verification_score",
    "overall score": "overall_score",
    "preferred contact": "preferred_contact_method",
    "preferred contact method": "preferred_contact_method",
    "tags": "tags",
    "source": "source",
    "source url": "source_url",
    "discovery source url": "discovery_source_url",
    "discovered at": "discovered_at",
    "last verified at": "last_verified_at",
    "outreach status": "outreach_status",
    "owner": "owner",
    "do not contact": "do_not_contact",
    "status": "status",
    "notes": "notes",
    "value proposition": "value_proposition",
}


def _normalize_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").replace("-", " ").split())


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, list):
        return ", ".join(str(item) for item in value if item not in (None, ""))
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, str) and value.startswith(("=", "+", "-")):
        return f"'{value}"
    return value


def _row_to_identity_lead(headers: List[str], values: Iterable[Any]) -> Dict[str, Any]:
    row = {_HEADER_TO_FIELD.get(header, header.replace(" ", "_")): value for header, value in zip(headers, values)}
    if not row.get("youtube_channel") and row.get("youtube_url"):
        row["youtube_channel"] = normalize_social_handle(str(row["youtube_url"]), "youtube")
    return row


def inspect_vento_workbook(path: str) -> Tuple[str, List[str], int, Set[str]]:
    """Return target sheet, normalized headers, old row count, and dedup keys."""
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        worksheet = workbook["Vento Influencers"] if "Vento Influencers" in workbook.sheetnames else workbook.active
        raw_headers = [str(cell.value or "").strip() for cell in worksheet[1]]
        normalized_headers = [_normalize_header(header) for header in raw_headers]
        recognized = sum(1 for header in normalized_headers if header in _HEADER_TO_FIELD)
        if recognized < 2 or not any(header in normalized_headers for header in ("creator name", "instagram handle", "tiktok handle", "email")):
            raise ValueError("The selected worksheet is not a recognizable Vento influencer export")
        identity_keys: Set[str] = set()
        old_rows = 0
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            if not any(value not in (None, "") for value in values):
                continue
            old_rows += 1
            identity_keys.update(creator_identity_keys(_row_to_identity_lead(normalized_headers, values)))
        return worksheet.title, normalized_headers, old_rows, identity_keys
    finally:
        workbook.close()


def _read_csv(path: str) -> Tuple[bytes, str, str, Any, List[str], List[List[str]]]:
    raw_bytes = Path(path).read_bytes()
    decoded = ""
    encoding = "utf-8"
    for candidate in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            decoded = raw_bytes.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if not decoded and raw_bytes:
        raise ValueError("CSV encoding is not supported")
    try:
        dialect = csv.Sniffer().sniff(decoded[:4000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(decoded, newline=""), dialect=dialect))
    if not rows:
        raise ValueError("The uploaded CSV is empty")
    raw_headers = [str(value or "").strip() for value in rows[0]]
    normalized_headers = [_normalize_header(header) for header in raw_headers]
    return raw_bytes, decoded, encoding, dialect, normalized_headers, rows[1:]


def inspect_vento_csv(path: str) -> Tuple[str, List[str], int, Set[str]]:
    """Read an old CSV without rewriting it and collect historical identities."""
    _, _, _, _, normalized_headers, rows = _read_csv(path)
    recognized = sum(1 for header in normalized_headers if header in _HEADER_TO_FIELD)
    if recognized < 2 or not any(header in normalized_headers for header in ("creator name", "name", "instagram handle", "instagram", "tiktok handle", "tiktok", "email")):
        raise ValueError("The uploaded CSV is not a recognizable influencer lead export")
    identity_keys: Set[str] = set()
    old_rows = 0
    for values in rows:
        if not any(str(value or "").strip() for value in values):
            continue
        old_rows += 1
        identity_keys.update(creator_identity_keys(_row_to_identity_lead(normalized_headers, values)))
    return "CSV", normalized_headers, old_rows, identity_keys


def _copy_row_style(worksheet: Any, source_row: int, target_row: int, column_count: int) -> None:
    if source_row < 2:
        return
    for column in range(1, column_count + 1):
        source = worksheet.cell(source_row, column)
        target = worksheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.protection:
            target.protection = copy(source.protection)
    if source_row in worksheet.row_dimensions:
        worksheet.row_dimensions[target_row].height = worksheet.row_dimensions[source_row].height


def append_leads_to_workbook(input_path: str, output_path: str, leads: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Preserve the old workbook and append unique new lead rows to a new copy."""
    workbook = load_workbook(input_path, data_only=False)
    try:
        worksheet = workbook["Vento Influencers"] if "Vento Influencers" in workbook.sheetnames else workbook.active
        raw_headers = [str(cell.value or "").strip() for cell in worksheet[1]]
        normalized_headers = [_normalize_header(header) for header in raw_headers]
        recognized = sum(1 for header in normalized_headers if header in _HEADER_TO_FIELD)
        if recognized < 2:
            raise ValueError("The selected worksheet is not a recognizable Vento influencer export")

        original_last_row = worksheet.max_row
        appended = 0
        for lead in leads:
            row_number = worksheet.max_row + 1
            _copy_row_style(worksheet, original_last_row, row_number, len(raw_headers))
            for column, normalized_header in enumerate(normalized_headers, start=1):
                field = _HEADER_TO_FIELD.get(normalized_header)
                value = lead.get(field) if field else ""
                worksheet.cell(row_number, column, _cell_value(value))
            appended += 1

        if appended:
            final_row = worksheet.max_row
            final_column = get_column_letter(max(1, len(raw_headers)))
            if worksheet.auto_filter.ref:
                worksheet.auto_filter.ref = f"A1:{final_column}{final_row}"
            for table in worksheet.tables.values():
                table.ref = f"A1:{final_column}{final_row}"

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output_path)
        return {
            "sheet_name": worksheet.title,
            "old_rows": max(0, original_last_row - 1),
            "new_rows": appended,
            "total_rows": max(0, worksheet.max_row - 1),
        }
    finally:
        workbook.close()


def append_leads_to_csv(input_path: str, output_path: str, leads: List[Dict[str, Any]]) -> Dict[str, Any]:
    """Keep the original CSV bytes intact and append new rows using its dialect."""
    raw_bytes, decoded, encoding, dialect, normalized_headers, rows = _read_csv(input_path)
    recognized = sum(1 for header in normalized_headers if header in _HEADER_TO_FIELD)
    if recognized < 2:
        raise ValueError("The uploaded CSV is not a recognizable influencer lead export")
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, dialect=dialect)
    for lead in leads:
        writer.writerow(
            [
                _cell_value(lead.get(_HEADER_TO_FIELD.get(header))) if _HEADER_TO_FIELD.get(header) else ""
                for header in normalized_headers
            ]
        )
    append_text = buffer.getvalue()
    separator = b"" if not raw_bytes or raw_bytes.endswith((b"\n", b"\r")) or not append_text else b"\r\n"
    append_encoding = "utf-8" if encoding == "utf-8-sig" else encoding
    Path(output_path).write_bytes(raw_bytes + separator + append_text.encode(append_encoding, errors="replace"))
    old_rows = sum(1 for row in rows if any(str(value or "").strip() for value in row))
    return {"sheet_name": "CSV", "old_rows": old_rows, "new_rows": len(leads), "total_rows": old_rows + len(leads)}


def inspect_vento_file(path: str) -> Tuple[str, List[str], int, Set[str]]:
    if Path(path).suffix.lower() == ".csv":
        return inspect_vento_csv(path)
    return inspect_vento_workbook(path)


def append_leads_to_file(input_path: str, output_path: str, leads: List[Dict[str, Any]]) -> Dict[str, Any]:
    if Path(input_path).suffix.lower() == ".csv":
        return append_leads_to_csv(input_path, output_path, leads)
    return append_leads_to_workbook(input_path, output_path, leads)
