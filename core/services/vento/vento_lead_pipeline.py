"""Production-ready implementation of the Vento Plan 3 two-pass lead finder."""
from __future__ import annotations

import asyncio
import csv
import io
import ipaddress
import json
import logging
import os
import re
import socket
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set, Tuple
from urllib.parse import urljoin, urlparse

import httpx
from openpyxl import load_workbook

from core.services.vento.apify_enrichment import ApifyCreatorEnricher
from core.services.vento.discovery import (
    EMAIL_RE,
    VALID_CATEGORIES,
    VENTO_MAX_ENRICHMENT_QUERIES,
    VENTO_RESULTS_PER_DISCOVERY_QUERY,
    VENTO_RESULTS_PER_ENRICHMENT_QUERY,
    VENTO_SEARCH_DEPTH,
    build_discovery_queries,
    build_enrichment_query,
    extract_emails,
    extract_phones,
    extract_social_profile_name,
    extract_social_handles,
    extract_websites,
    infer_name_from_title,
    is_social_or_directory_url,
    normalize_social_handle,
    normalized_domain,
)
from core.services.vento.verification import parse_follower_count
from core.tools.external.html_search_fallback import HTMLSearchFallbackClient
from core.tools.external.tavily_client import TavilyClient

try:
    from google import genai
    from google.genai import types as genai_types
except ImportError:  # pragma: no cover - dependencies are optional outside production
    genai = None
    genai_types = None


_CATEGORY_NICHES = {
    "dog_parent_influencers": "Dog Parent Influencer",
    "online_pet_businesses": "Online Pet Business",
    "pet_shops": "Pet Shop",
    "pet_grooming": "Pet Grooming",
}
_PET_TERMS = frozenset(
    {
        "pet", "pets", "dog", "dogs", "puppy", "puppies", "canine", "groomer",
        "grooming", "animal", "paw", "paws", "veterinary", "trainer", "kennel",
    }
)
_CLEARLY_UNRELATED_TERMS = frozenset(
    {"car dealership", "real estate", "mortgage", "dentist", "plumber", "law firm", "insurance agency"}
)
_PLACEHOLDER_NAME_KEYS = frozenset(
    {
        "username",
        "yourusername",
        "instagramusername",
        "tiktokusername",
        "user",
        "name",
        "profile",
        "account",
        "creator",
        "influencer",
        "unknown",
        "na",
        "notavailable",
        "example",
        "exampleuser",
        "testuser",
    }
)
_CONTACT_LINK_RE = re.compile(
    r"href=[\"']([^\"']*(?:contact|about|media-kit|press|collaborat|partnership|work-with)[^\"']*)[\"']",
    re.I,
)
_MAILTO_RE = re.compile(r"mailto:([^\"'?#\s>]+)", re.I)
_TEL_RE = re.compile(r"tel:([^\"'?#\s>]+)", re.I)
_JSON_PHONE_RE = re.compile(r'[\"\']telephone[\"\']\s*:\s*[\"\']([^\"\']+)', re.I)
_SCRIPT_STYLE_RE = re.compile(r"<(?:script|style|noscript|svg)\b[^>]*>.*?</(?:script|style|noscript|svg)>", re.I | re.S)
_HTML_TAG_RE = re.compile(r"<[^>]+>")
_WHITESPACE_RE = re.compile(r"\s+")
_FOLLOWER_COUNT_RE = re.compile(r"(?P<count>\d[\d,.]*\s*[KMB]?\+?)\s+(?:followers?|fans?)\b", re.I)
_FOLLOWER_COUNT_REVERSE = re.compile(r"\b(?:followers?|fans?)\s*[:=-]?\s*(?P<count>\d[\d,.]*\s*[KMB]?\+?)", re.I)
_BAD_CREATOR_EMAIL_LOCALS = frozenset(
    {
        "admin",
        "ads",
        "advertise",
        "advertising",
        "careers",
        "circulation",
        "corrections",
        "editor",
        "editorial",
        "help",
        "jobs",
        "legal",
        "news",
        "newsroom",
        "no-reply",
        "noreply",
        "press",
        "privacy",
        "subscribe",
        "subscriptions",
        "support",
        "tips",
        "webmaster",
    }
)
_BAD_CREATOR_EMAIL_DOMAINS = frozenset(
    {
        "abc15.com",
        "apify.com",
        "cbs.com",
        "cbsnews.com",
        "collabstr.com",
        "denverlifemagazine.com",
        "denverpost.com",
        "facebook.com",
        "favikon.com",
        "feedspot.com",
        "greenwichtime.com",
        "heepsy.com",
        "influencers.club",
        "instagram.com",
        "insense.pro",
        "meta.com",
        "modash.io",
        "patch.com",
        "petfinder.com",
        "popularpays.com",
        "scni.com",
        "sfgate.com",
        "socialcat.com",
        "tiktok.com",
    }
)
_CREATOR_EMAIL_HINT_LOCALS = frozenset(
    {
        "booking",
        "bookings",
        "brand",
        "brands",
        "business",
        "collab",
        "collabs",
        "collabwith",
        "contact",
        "hello",
        "hi",
        "inquiries",
        "management",
        "manager",
        "media",
        "mgmt",
        "partnership",
        "partnerships",
        "pr",
        "team",
        "work",
    }
)
_PERSONAL_EMAIL_DOMAINS = frozenset(
    {
        "aol.com",
        "gmail.com",
        "hotmail.com",
        "icloud.com",
        "live.com",
        "me.com",
        "outlook.com",
        "proton.me",
        "protonmail.com",
        "yahoo.com",
    }
)
_BAD_EMAIL_CONTEXT_TERMS = frozenset(
    {
        "advertise",
        "advertising",
        "circulation",
        "editor",
        "editorial",
        "newsroom",
        "press release",
        "publisher",
        "reporter",
        "subscribe",
    }
)


class VentoLeadPipeline:
    """Discover, enrich, deduplicate, score, and gate pet-industry leads."""

    def __init__(
        self,
        *,
        tavily_client: Any = None,
        fallback_search_client: Any = None,
        gemini_client: Any = None,
        gemini_api_key: Optional[str] = None,
        gemini_model: Optional[str] = None,
        apify_enricher: Any = None,
        social_provider: Any = None,
        verifier: Any = None,
    ) -> None:
        self.logger = logging.getLogger("agent.runtime")
        self.tavily = tavily_client
        if self.tavily is None and os.getenv("TAVILY_API_KEY", "").strip():
            self.tavily = TavilyClient()
        self.fallback_search = fallback_search_client or HTMLSearchFallbackClient()

        api_key = gemini_api_key or os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
        self.gemini = gemini_client
        if self.gemini is None and api_key and genai is not None:
            self.gemini = genai.Client(api_key=api_key)
        self.gemini_model = (
            gemini_model
            or os.getenv("VENTO_GEMINI_MODEL", "").strip()
            or os.getenv("GEMINI_MODEL", "").strip()
            or "gemini-3.1-flash-lite"
        )
        self.gemini_fallback_models = list(
            dict.fromkeys([self.gemini_model, "gemini-3.1-flash-lite", "gemini-3.5-flash"])
        )

        self.discovery_concurrency = self._env_int("VENTO_DISCOVERY_CONCURRENCY", 3, 1, 6)
        self.enrichment_concurrency = self._env_int("VENTO_ENRICHMENT_CONCURRENCY", 5, 1, 10)
        self.scrape_concurrency = self._env_int("VENTO_SCRAPE_CONCURRENCY", 6, 1, 12)
        self.scrape_timeout = self._env_float("VENTO_SCRAPE_TIMEOUT_SECONDS", 8.0, 3.0, 20.0)
        self.search_timeout = self._env_float("VENTO_SEARCH_TIMEOUT_SECONDS", 12.0, 3.0, 45.0)
        self.discovery_timeout = self._env_float("VENTO_DISCOVERY_TIMEOUT_SECONDS", 45.0, 10.0, 180.0)
        self.enrichment_timeout = self._env_float("VENTO_ENRICHMENT_TIMEOUT_SECONDS", 90.0, 20.0, 240.0)
        self.apify_stage_timeout = self._env_float("VENTO_APIFY_STAGE_TIMEOUT_SECONDS", 70.0, 15.0, 240.0)
        self.scoring_timeout = self._env_float("VENTO_SCORING_TIMEOUT_SECONDS", 25.0, 5.0, 120.0)
        self.discovery_results_per_query = self._env_int(
            "VENTO_RESULTS_PER_DISCOVERY_QUERY",
            VENTO_RESULTS_PER_DISCOVERY_QUERY,
            5,
            20,
        )
        self.max_enrichment_candidates = self._env_int("VENTO_MAX_ENRICHMENT_CANDIDATES", 60, 10, 200)
        self.discovery_expansion_enabled = self._env_bool("VENTO_DISCOVERY_EXPANSION_ENABLED", True)
        self.max_discovery_expansion_queries = self._env_int("VENTO_MAX_DISCOVERY_EXPANSION_QUERIES", 6, 0, 12)
        self.apify_candidate_pool = self._env_int("VENTO_APIFY_CANDIDATE_POOL", 20, 1, 200)
        self.gemini_budget = self._env_int("VENTO_GEMINI_MAX_CALLS_PER_RUN", 40, 0, 40)
        self.max_upload_rows = self._env_int("VENTO_MAX_IMPORT_ROWS", 50_000, 1, 200_000)
        self.min_social_followers = self._env_int("VENTO_MIN_SOCIAL_FOLLOWERS", 50_000, 0, 10_000_000)
        self.max_social_followers = self._env_int("VENTO_MAX_SOCIAL_FOLLOWERS", 500_000, 0, 100_000_000)
        self.require_social_followers = self._env_bool("VENTO_REQUIRE_SOCIAL_FOLLOWERS", False)
        self.allow_above_max_social_followers = self._env_bool("VENTO_ALLOW_ABOVE_MAX_SOCIAL_FOLLOWERS", True)
        self.apify_enrichment_enabled = self._env_bool("VENTO_APIFY_ENRICHMENT_ENABLED", True)
        self.apify_enricher = apify_enricher if apify_enricher is not None else ApifyCreatorEnricher()
        self.social_profile_fetch_enabled = self._env_bool(
            "VENTO_SOCIAL_NAME_FETCH_ENABLED",
            not (self.apify_enrichment_enabled and getattr(self.apify_enricher, "configured", False)),
        )
        self._stats: Dict[str, Any] = self._empty_stats()

    @staticmethod
    def _empty_stats() -> Dict[str, int]:
        return {
            "tavily_calls": 0,
            "fallback_search_calls": 0,
            "search_failures": 0,
            "search_timeouts": 0,
            "discovery_expansion_queries": 0,
            "discovery_expansion_candidates": 0,
            "discovery_timeouts": 0,
            "enrichment_candidates_selected": 0,
            "enrichment_timeouts": 0,
            "scrape_requests": 0,
            "scrape_failures": 0,
            "social_name_requests": 0,
            "social_names_recovered": 0,
            "social_follower_requests": 0,
            "social_followers_recovered": 0,
            "apify_enabled": 0,
            "apify_prequalified_candidates": 0,
            "apify_candidates_skipped_prefilter": 0,
            "apify_profile_requests": 0,
            "apify_profiles_enriched": 0,
            "apify_cache_hits": 0,
            "apify_cache_misses": 0,
            "apify_profiles_skipped_by_budget": 0,
            "apify_followers_found": 0,
            "apify_emails_found": 0,
            "apify_platform_mismatches_dropped": 0,
            "apify_failures": 0,
            "apify_timeouts": 0,
            "gemini_calls": 0,
            "gemini_api_attempts": 0,
            "gemini_failures": 0,
            "scoring_timeouts": 0,
        }

    @staticmethod
    def _env_int(name: str, default: int, minimum: int, maximum: int) -> int:
        try:
            return max(minimum, min(maximum, int(os.getenv(name, str(default)))))
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _env_float(name: str, default: float, minimum: float, maximum: float) -> float:
        try:
            return max(minimum, min(maximum, float(os.getenv(name, str(default)))))
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _env_bool(name: str, default: bool = False) -> bool:
        value = os.getenv(name)
        if value is None:
            return default
        return value.strip().lower() in {"1", "true", "yes", "on"}

    async def _run_with_time_budget(
        self,
        coroutines: Sequence[Any],
        *,
        timeout: float,
        timeout_stat: str,
        label: str,
    ) -> List[Any]:
        tasks = [asyncio.create_task(coro) for coro in coroutines]
        if not tasks:
            return []
        done, pending = await asyncio.wait(tasks, timeout=timeout)
        if pending:
            self._stats[timeout_stat] = int(self._stats.get(timeout_stat, 0)) + len(pending)
            self.logger.warning(
                "vento_stage_time_budget_exceeded",
                extra={"stage": label, "timeout": timeout, "pending": len(pending), "completed": len(done)},
            )
            for task in pending:
                task.cancel()
            await asyncio.gather(*pending, return_exceptions=True)
        output: List[Any] = []
        for task in tasks:
            if task not in done:
                continue
            try:
                output.append(task.result())
            except Exception as exc:
                self._stats["search_failures"] += 1
                self.logger.warning("vento_stage_task_failed", extra={"stage": label, "error": str(exc)})
        return output

    async def run(
        self,
        *,
        category: str,
        location: str,
        target_count: int = 50,
        old_leads_csv: Optional[str] = None,
        old_leads_file: Optional[str] = None,
    ) -> Dict[str, Any]:
        category = str(category or "").strip().lower()
        location = _WHITESPACE_RE.sub(" ", str(location or "")).strip()
        if category not in VALID_CATEGORIES:
            raise ValueError("category must be one of: " + ", ".join(sorted(VALID_CATEGORIES)))
        if not location:
            raise ValueError("location is required")
        try:
            target_count = int(target_count)
        except (TypeError, ValueError) as exc:
            raise ValueError("target_count must be a number") from exc
        if not 1 <= target_count <= 200:
            raise ValueError("target_count must be between 1 and 200")

        self._stats = self._empty_stats()
        raw_candidates = await self._discover_candidates(category, location, target_count)
        enriched = await self._enrich_candidates(raw_candidates, category, location, target_count)
        try:
            enriched = await asyncio.wait_for(
                self._enrich_with_apify_if_enabled(enriched, category),
                timeout=self.apify_stage_timeout,
            )
        except asyncio.TimeoutError:
            self._stats["apify_timeouts"] += 1
            self._stats["apify_failures"] += 1
            self.logger.warning(
                "vento_apify_stage_timeout",
                extra={"timeout": self.apify_stage_timeout, "candidate_count": len(enriched)},
            )

        old_keys: Set[str] = set()
        old_row_count = 0
        if old_leads_csv:
            old_keys = self._extract_dedup_keys_from_csv(old_leads_csv)
            old_row_count = self._count_csv_rows(old_leads_csv)
        elif old_leads_file:
            old_rows = self._read_old_lead_rows(old_leads_file)
            old_row_count = len(old_rows)
            old_keys = self._extract_dedup_keys_from_rows(old_rows)

        deduped, duplicate_count = self._dedup(enriched, old_keys)
        try:
            scored = await asyncio.wait_for(
                self._score_leads(deduped, category, location),
                timeout=self.scoring_timeout,
            )
        except asyncio.TimeoutError:
            self._stats["scoring_timeouts"] += 1
            self._stats["gemini_failures"] += 1
            self.logger.warning(
                "vento_scoring_timeout",
                extra={"timeout": self.scoring_timeout, "lead_count": len(deduped)},
            )
            scored = deduped
            for lead in scored:
                lead["relevance_score"] = self._heuristic_score(lead, category, location)
                lead["scoring_method"] = "heuristic_timeout_fallback"
        usable, rejected = self._quality_gate(scored, category)
        usable.sort(
            key=lambda lead: (
                self._follower_sort_bucket(lead, category),
                -float(lead.get("relevance_score") or 0),
                -self._contact_count(lead, category),
                self._lead_name(lead).lower(),
            )
        )
        final = usable[:target_count]

        return {
            "leads": final,
            "rejected_leads": rejected,
            "category": category,
            "location": location,
            "target_count": target_count,
            "returned_count": len(final),
            "shortfall": max(0, target_count - len(final)),
            "metrics": {
                "raw_discovered": len(raw_candidates),
                "enriched": len(enriched),
                "duplicates_skipped": duplicate_count,
                "usable": len(usable),
                "rejected": len(rejected),
                "old_rows": old_row_count,
                "with_email": sum(bool(lead.get("email")) for lead in final),
                "with_phone": sum(bool(lead.get("phone")) for lead in final),
                "with_instagram": sum(bool(lead.get("instagram_handle")) for lead in final),
                "with_tiktok": sum(bool(lead.get("tiktok_handle")) for lead in final),
                "with_website": sum(bool(lead.get("website")) for lead in final),
                "with_followers": sum(parse_follower_count(lead.get("follower_count")) is not None for lead in final),
                "followers_in_target_range": sum(self._follower_sort_bucket(lead, category) == 0 for lead in final) if category == "dog_parent_influencers" else 0,
                "followers_above_target_range": sum(self._follower_sort_bucket(lead, category) == 2 for lead in final) if category == "dog_parent_influencers" else 0,
                "followers_unknown": sum(self._follower_sort_bucket(lead, category) == 3 for lead in final) if category == "dog_parent_influencers" else 0,
                "min_social_followers": self.min_social_followers if category == "dog_parent_influencers" else 0,
                "max_social_followers": self.max_social_followers if category == "dog_parent_influencers" else 0,
                "require_social_followers": self.require_social_followers if category == "dog_parent_influencers" else False,
                **self._stats,
            },
            "steps": [
                {"step": 1, "phase": "discovery", "count": len(raw_candidates)},
                {"step": 2, "phase": "enrichment_and_scrape", "count": len(enriched)},
                {"step": 3, "phase": "deduplication", "count": len(deduped)},
                {"step": 4, "phase": "scoring", "count": len(scored)},
                {"step": 5, "phase": "quality_gate", "count": len(final)},
            ],
        }

    def _build_discovery_queries(self, category: str, location: str) -> List[str]:
        return build_discovery_queries(category, location)

    async def _search(self, query: str, *, max_results: int) -> Tuple[Dict[str, Any], str]:
        payload = {
            "query": query,
            "search_depth": VENTO_SEARCH_DEPTH,
            "max_results": max_results,
            "include_raw_content": False,
        }
        if self.tavily is not None:
            self._stats["tavily_calls"] += 1
            try:
                return await asyncio.wait_for(self.tavily.search(payload), timeout=self.search_timeout), "tavily"
            except asyncio.TimeoutError:
                self._stats["search_timeouts"] += 1
                self._stats["search_failures"] += 1
                self.logger.warning("vento_tavily_search_timeout", extra={"query": query, "timeout": self.search_timeout})
            except Exception as exc:
                self._stats["search_failures"] += 1
                self.logger.warning("vento_tavily_search_failed", extra={"query": query, "error": str(exc)})
        self._stats["fallback_search_calls"] += 1
        try:
            return await asyncio.wait_for(self.fallback_search.search(payload), timeout=self.search_timeout), "web_crawler_fallback"
        except asyncio.TimeoutError:
            self._stats["search_timeouts"] += 1
            self._stats["search_failures"] += 1
            self.logger.warning("vento_fallback_search_timeout", extra={"query": query, "timeout": self.search_timeout})
            return {"results": []}, "unavailable"
        except Exception as exc:
            self._stats["search_failures"] += 1
            self.logger.warning("vento_fallback_search_failed", extra={"query": query, "error": str(exc)})
            return {"results": []}, "unavailable"

    async def _discover_candidates(self, category: str, location: str, target_count: int) -> List[Dict[str, Any]]:
        queries = self._build_discovery_queries(category, location)
        semaphore = asyncio.Semaphore(self.discovery_concurrency)

        async def search_one(query: str) -> Tuple[str, Dict[str, Any], str]:
            async with semaphore:
                data, provider = await self._search(query, max_results=self.discovery_results_per_query)
                return query, data, provider

        candidates: List[Dict[str, Any]] = []
        seen = set()

        def add_batches(batches: Sequence[Tuple[str, Dict[str, Any], str]]) -> int:
            added = 0
            for query, data, provider in batches:
                before = len(candidates)
                self._append_discovery_candidates(
                    candidates,
                    seen,
                    query=query,
                    data=data,
                    provider=provider,
                    category=category,
                    location=location,
                )
                added += len(candidates) - before
            return added

        batches = await self._run_with_time_budget(
            [search_one(query) for query in queries],
            timeout=self.discovery_timeout,
            timeout_stat="discovery_timeouts",
            label="discovery",
        )
        add_batches(batches)
        desired_raw = min(max(target_count * 3, 60), 140) if category == "dog_parent_influencers" else 60
        if (
            self.discovery_expansion_enabled
            and self.max_discovery_expansion_queries > 0
            and category == "dog_parent_influencers"
            and len(candidates) < min(desired_raw, target_count * 2)
        ):
            expansion_queries = self._build_discovery_expansion_queries(location)[: self.max_discovery_expansion_queries]
            self._stats["discovery_expansion_queries"] += len(expansion_queries)
            expansion_batches = await self._run_with_time_budget(
                [search_one(query) for query in expansion_queries],
                timeout=max(10.0, self.discovery_timeout / 2),
                timeout_stat="discovery_timeouts",
                label="discovery_expansion",
            )
            self._stats["discovery_expansion_candidates"] += add_batches(expansion_batches)

        candidates.sort(key=lambda lead: (-self._candidate_priority(lead, category), self._lead_name(lead).lower()))
        return candidates[:desired_raw]

    def _append_discovery_candidates(
        self,
        candidates: List[Dict[str, Any]],
        seen: Set[Tuple[str, str]],
        *,
        query: str,
        data: Dict[str, Any],
        provider: str,
        category: str,
        location: str,
    ) -> None:
        for item in data.get("results") or []:
            if not isinstance(item, dict):
                continue
            url = str(item.get("url") or "").strip()
            title = str(item.get("title") or "").strip()
            content = str(item.get("content") or "").strip()
            combined = f"{title} {content} {url}"
            socials = extract_social_handles(combined, url)
            social_name = ""
            social_platform = "instagram" if socials["instagram"] else "tiktok" if socials["tiktok"] else ""
            social_handle = (socials.get(social_platform) or [""])[0] if social_platform else ""
            if social_platform:
                social_name = extract_social_profile_name(combined, social_platform, social_handle)
            name = social_name or infer_name_from_title(title, url)
            name_source = "social_search_metadata" if social_name else "search_result_title"
            social_post = bool(
                social_platform == "tiktok" and "/video/" in url.lower()
                or social_platform == "instagram"
                and re.search(r"instagram\.com/(?:p|reel|reels|tv|stories)/", url, re.I)
            )
            if not name or (social_post and not social_name):
                name = social_handle.replace("_", " ").replace(".", " ").title() if social_handle else ""
                name_source = "handle_fallback"
            identity = (normalized_domain(url, allow_generic=True), self._normalize_name(name))
            if not name or self._looks_like_placeholder_name(name) or identity in seen:
                continue
            seen.add(identity)
            websites = [] if is_social_or_directory_url(url) else [url]
            evidence = f"{query} {combined}"
            raw_emails = extract_emails(combined)
            creator_emails = (
                self._creator_emails(raw_emails, combined, name=name, socials=socials)
                if category == "dog_parent_influencers"
                else raw_emails
            )
            candidate = {
                self._name_field(category): name,
                "email": (creator_emails or [""])[0],
                "phone": (extract_phones(combined) or [""])[0],
                "website": websites[0] if websites else "",
                "instagram_handle": self._format_handle((socials["instagram"] or [""])[0]),
                "tiktok_handle": self._format_handle((socials["tiktok"] or [""])[0]),
                "location": location if self._location_matches(evidence, location) else "",
                "niche": _CATEGORY_NICHES[category],
                "source": "tavily_discovery",
                "category": category,
                "source_url": url,
                "creator_name_source": name_source if category == "dog_parent_influencers" else "search_result_title",
                "creator_name_confidence": (
                    "high" if name_source == "social_search_metadata"
                    else "low" if name_source == "handle_fallback"
                    else "medium"
                ) if category == "dog_parent_influencers" else "medium",
                "discovery_provider": provider,
                "discovery_query": query,
                "evidence": evidence[:8_000],
            }
            if category == "dog_parent_influencers":
                self._record_follower_evidence(candidate, combined, source_url=url, status="search_observed")
            candidates.append(candidate)

    @staticmethod
    def _build_discovery_expansion_queries(location: str) -> List[str]:
        return [
            f'"dog mom" "{location}" instagram tiktok creator',
            f'"dog dad" "{location}" instagram tiktok creator',
            f'"dog influencer" "{location}" instagram',
            f'"dog influencer" "{location}" tiktok',
            f'"pet influencer" "{location}" instagram tiktok followers',
            f'"pet content creator" "{location}" instagram',
            f'"pet content creator" "{location}" tiktok',
            f'"dog parent" "{location}" "instagram.com"',
            f'"dog parent" "{location}" "tiktok.com/@"',
            f'"dog account" "{location}" instagram',
            f'"dog trainer" "{location}" instagram creator',
            f'"dog travel" "{location}" instagram tiktok',
        ]

    async def _enrich_candidates(
        self,
        candidates: Sequence[Dict[str, Any]],
        category: str,
        location: str,
        target_count: int,
    ) -> List[Dict[str, Any]]:
        selected_limit = min(len(candidates), self.max_enrichment_candidates)
        if category == "dog_parent_influencers":
            selected_limit = min(selected_limit, max(target_count * 2, target_count + 10))
        selected = list(candidates[:selected_limit])
        semaphore = asyncio.Semaphore(self.enrichment_concurrency)

        async def enrich(candidate: Dict[str, Any]) -> Dict[str, Any]:
            async with semaphore:
                return await self._enrich_single_candidate(candidate, category, location)

        self._stats["enrichment_candidates_selected"] = len(selected)
        return list(
            await self._run_with_time_budget(
                [enrich(candidate) for candidate in selected],
                timeout=self.enrichment_timeout,
                timeout_stat="enrichment_timeouts",
                label="enrichment",
            )
        )

    async def _enrich_with_apify_if_enabled(
        self,
        leads: Sequence[Dict[str, Any]],
        category: str,
    ) -> List[Dict[str, Any]]:
        if category != "dog_parent_influencers":
            return [dict(lead) for lead in leads]
        if not self.apify_enrichment_enabled:
            return [dict(lead) for lead in leads]
        if not getattr(self.apify_enricher, "configured", False):
            return [dict(lead) for lead in leads]
        shortlisted, skipped = self._shortlist_for_apify(leads)
        self._stats["apify_prequalified_candidates"] += len(shortlisted)
        self._stats["apify_candidates_skipped_prefilter"] += skipped
        if not shortlisted:
            return [dict(lead) for lead in leads]
        try:
            enriched_shortlist, stats = await self.apify_enricher.enrich_profiles(shortlisted)
        except Exception as exc:
            self._stats["apify_failures"] += 1
            self.logger.warning("vento_apify_enrichment_failed", extra={"error": str(exc)})
            return [dict(lead) for lead in leads]
        for key, value in (stats or {}).items():
            if isinstance(value, int):
                self._stats[key] = int(self._stats.get(key, 0)) + value
        return self._merge_apify_shortlist(leads, enriched_shortlist)

    def _shortlist_for_apify(self, leads: Sequence[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], int]:
        candidates = [dict(lead) for lead in leads if self._apify_prefilter_passes(lead)]
        candidates.sort(key=lambda lead: (-self._pre_apify_score(lead), self._lead_name(lead).lower()))
        selected = candidates[:self.apify_candidate_pool]
        return selected, max(0, len(leads) - len(selected))

    def _apify_prefilter_passes(self, lead: Dict[str, Any]) -> bool:
        if not self._lead_name(lead) or self._looks_like_placeholder_name(self._lead_name(lead)):
            return False
        if self._contact_count(lead, "dog_parent_influencers") < 1:
            return False
        if self._clearly_not_pet_related(lead):
            return False
        count = parse_follower_count(lead.get("follower_count"))
        if count is not None and count < self.min_social_followers:
            return False
        evidence = str(lead.get("evidence") or "").lower()
        return any(term in evidence for term in _PET_TERMS) or bool(lead.get("instagram_handle") or lead.get("tiktok_handle"))

    def _pre_apify_score(self, lead: Dict[str, Any]) -> int:
        evidence = str(lead.get("evidence") or "").lower()
        score = 0
        score += self._contact_count(lead, "dog_parent_influencers") * 8
        score += sum(term in evidence for term in _PET_TERMS) * 4
        if lead.get("email"):
            score += 10
        if lead.get("instagram_handle") and lead.get("tiktok_handle"):
            score += 10
        if lead.get("location"):
            score += 6
        count = parse_follower_count(lead.get("follower_count"))
        if count is not None:
            score += 8
            if count >= self.min_social_followers and (self.max_social_followers <= 0 or count <= self.max_social_followers):
                score += 15
            elif count > self.max_social_followers > 0:
                score += 4
        if lead.get("creator_name_source") == "social_search_metadata":
            score += 5
        elif lead.get("creator_name_source") == "handle_fallback":
            score -= 3
        return score

    def _merge_apify_shortlist(
        self,
        original: Sequence[Dict[str, Any]],
        enriched_shortlist: Sequence[Dict[str, Any]],
    ) -> List[Dict[str, Any]]:
        enriched_by_key: Dict[str, Dict[str, Any]] = {}
        for lead in enriched_shortlist:
            key = self._apify_merge_key(lead)
            if key:
                enriched_by_key[key] = dict(lead)
        output: List[Dict[str, Any]] = []
        for lead in original:
            key = self._apify_merge_key(lead)
            output.append(enriched_by_key.get(key, dict(lead)))
        return output

    @classmethod
    def _apify_merge_key(cls, lead: Dict[str, Any]) -> str:
        instagram = str(lead.get("instagram_handle") or "").strip().lower().lstrip("@")
        tiktok = str(lead.get("tiktok_handle") or "").strip().lower().lstrip("@")
        if instagram:
            return f"instagram:{instagram}"
        if tiktok:
            return f"tiktok:{tiktok}"
        return f"name:{cls._normalize_name(cls._lead_name(lead))}"

    async def _enrich_single_candidate(
        self,
        candidate: Dict[str, Any],
        category: str,
        location: str,
    ) -> Dict[str, Any]:
        lead = dict(candidate)
        query = build_enrichment_query(category, self._lead_name(lead), location)
        data, provider = await self._search(query, max_results=VENTO_RESULTS_PER_ENRICHMENT_QUERY)
        results = [item for item in (data.get("results") or []) if isinstance(item, dict)]
        matched_results = self._identity_matched_results(lead, results)
        blocks = []
        result_urls = []
        for item in matched_results:
            item_url = str(item.get("url") or "").strip()
            result_urls.append(item_url)
            blocks.append(f"{item.get('title') or ''} {item.get('content') or ''} {item_url}")
        combined = " ".join(blocks)
        evidence = f"{lead.get('evidence') or ''} {query} {combined}"
        emails = extract_emails(combined)
        phones = extract_phones(combined)
        socials = extract_social_handles(combined, *result_urls)
        websites = extract_websites(combined, *result_urls)

        if category == "dog_parent_influencers":
            merged_socials = self._lead_socials(lead, socials)
            creator_emails = self._creator_emails(emails, combined, name=self._lead_name(lead), socials=merged_socials)
            current_email = str(lead.get("email") or "").strip().lower()
            if not current_email or not self._is_creator_email(
                current_email,
                str(lead.get("evidence") or ""),
                name=self._lead_name(lead),
                socials=merged_socials,
            ):
                lead["email"] = creator_emails[0] if creator_emails else ""
        else:
            lead["email"] = lead.get("email") or (emails[0] if emails else "")
        lead["phone"] = lead.get("phone") or (phones[0] if phones else "")
        lead["instagram_handle"] = lead.get("instagram_handle") or self._format_handle((socials["instagram"] or [""])[0])
        lead["tiktok_handle"] = lead.get("tiktok_handle") or self._format_handle((socials["tiktok"] or [""])[0])
        lead["website"] = lead.get("website") or self._best_website(websites, self._lead_name(lead))
        lead["location"] = lead.get("location") or (location if self._location_matches(evidence, location) else "")
        lead["enrichment_provider"] = provider
        lead["enrichment_query"] = query
        lead["evidence"] = evidence[:20_000]
        lead["instagram_url"] = self._social_url(lead.get("instagram_handle"), "instagram")
        lead["tiktok_url"] = self._social_url(lead.get("tiktok_handle"), "tiktok")
        if category == "dog_parent_influencers":
            self._record_follower_evidence(
                lead,
                combined,
                source_url=(result_urls[0] if result_urls else str(lead.get("source_url") or "")),
                status="search_observed",
            )

        for platform in ("instagram", "tiktok"):
            handle = str(lead.get(f"{platform}_handle") or "").strip()
            if not handle:
                continue
            for item in matched_results:
                indexed_name = extract_social_profile_name(
                    f"{item.get('title') or ''} {item.get('content') or ''}", platform, handle
                )
                if indexed_name:
                    lead["creator_name"] = indexed_name
                    lead["creator_name_source"] = "social_search_metadata"
                    lead["creator_name_confidence"] = "medium"
                    break
            if lead.get("creator_name_source") == "social_search_metadata":
                break

        if lead.get("website"):
            scraped = await self._scrape_candidate_website(str(lead["website"]))
            # The candidate's own site is more authoritative than snippets. Its
            # contacts replace snippet contacts to prevent cross-business mixing.
            if scraped["emails"]:
                website_email = self._best_email_for_website(
                    scraped["emails"],
                    str(lead["website"]),
                    lead=lead if category == "dog_parent_influencers" else None,
                )
                if website_email or category != "dog_parent_influencers":
                    lead["email"] = website_email
            if scraped["phones"]:
                lead["phone"] = scraped["phones"][0]
            lead["instagram_handle"] = lead.get("instagram_handle") or self._format_handle((scraped["socials"]["instagram"] or [""])[0])
            lead["tiktok_handle"] = lead.get("tiktok_handle") or self._format_handle((scraped["socials"]["tiktok"] or [""])[0])
            lead["instagram_url"] = self._social_url(lead.get("instagram_handle"), "instagram")
            lead["tiktok_url"] = self._social_url(lead.get("tiktok_handle"), "tiktok")
            lead["scraped_urls"] = scraped["urls"]
            lead["evidence"] = f"{lead['evidence']} {scraped['text']}"[:20_000]
        if category == "dog_parent_influencers":
            await self._enrich_creator_name_from_social_profile(lead)
        return lead

    async def _enrich_creator_name_from_social_profile(self, lead: Dict[str, Any]) -> None:
        """Prefer public social metadata over inferred titles and recover follower counts."""
        if not self.social_profile_fetch_enabled:
            return
        name_recovered = False
        for platform in ("instagram", "tiktok"):
            handle = str(lead.get(f"{platform}_handle") or "").strip().lstrip("@")
            url = str(lead.get(f"{platform}_url") or "").strip()
            if not handle or not url:
                continue
            if not name_recovered:
                name = await self._fetch_social_profile_name(url, platform, handle)
                if name:
                    lead["creator_name"] = name
                    lead["creator_name_source"] = f"{platform}_profile_metadata"
                    lead["creator_name_confidence"] = "high"
                    lead["creator_name_profile_url"] = url
                    self._stats["social_names_recovered"] += 1
                    name_recovered = True
            if parse_follower_count(lead.get("follower_count")) is None:
                follower_count = await self._fetch_social_profile_follower_count(url, platform, handle)
                if follower_count is not None:
                    self._set_follower_count(
                        lead,
                        follower_count,
                        source_url=url,
                        status=f"{platform}_profile_metadata",
                    )

    async def _fetch_social_profile_name(self, url: str, platform: str, handle: str) -> str:
        if not await self._safe_public_url(url):
            return ""
        self._stats["social_name_requests"] += 1
        timeout_seconds = self._env_float("VENTO_SOCIAL_NAME_TIMEOUT_SECONDS", 6.0, 2.0, 15.0)
        timeout = httpx.Timeout(timeout_seconds, connect=min(4.0, timeout_seconds))
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
            ),
            "Accept-Language": "en-US,en;q=0.9",
        }
        try:
            async with httpx.AsyncClient(timeout=timeout, follow_redirects=True, headers=headers) as client:
                page = await self._fetch_html(client, url)
        except Exception:
            return ""
        if not page:
            return ""
        return extract_social_profile_name(page[1], platform, handle)

    async def _fetch_social_profile_follower_count(self, url: str, platform: str, handle: str) -> Optional[int]:
        if not await self._safe_public_url(url):
            return None
        self._stats["social_follower_requests"] += 1
        timeout_seconds = self._env_float("VENTO_SOCIAL_NAME_TIMEOUT_SECONDS", 6.0, 2.0, 15.0)
        timeout = httpx.Timeout(timeout_seconds, connect=min(4.0, timeout_seconds))
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"
            ),
            "Accept-Language": "en-US,en;q=0.9",
        }
        try:
            async with httpx.AsyncClient(timeout=timeout, follow_redirects=True, headers=headers) as client:
                page = await self._fetch_html(client, url)
        except Exception:
            return None
        if not page:
            return None
        body = page[1]
        normalized_handle = handle.lower().lstrip("@")
        if normalized_handle and normalized_handle not in body.lower():
            return None
        return self._extract_follower_count(body)

    def _identity_matched_results(
        self,
        candidate: Dict[str, Any],
        results: Sequence[Dict[str, Any]],
    ) -> List[Dict[str, Any]]:
        """Keep enrichment evidence tied to one candidate/official domain."""
        if not results:
            return []
        candidate_domain = normalized_domain(str(candidate.get("website") or ""), allow_generic=True)
        name_tokens = {
            token for token in re.findall(r"[a-z0-9]+", self._lead_name(candidate).lower())
            if len(token) >= 4 and token not in {"best", "premier", "austin", "texas", "mobile", "online", "shop", "store"}
        }

        def score(item: Dict[str, Any]) -> int:
            url = str(item.get("url") or "")
            domain = normalized_domain(url, allow_generic=True)
            title = str(item.get("title") or "").lower()
            content = str(item.get("content") or "").lower()
            haystack = f"{title} {content} {domain}"
            domain_match = bool(candidate_domain and domain == candidate_domain)
            token_hits = sum(token in haystack for token in name_tokens)
            return (20 if domain_match else 0) + token_hits * 2 + (1 if not is_social_or_directory_url(url) else 0)

        ranked = sorted(results, key=score, reverse=True)
        best = ranked[0]
        best_score = score(best)
        best_domain = normalized_domain(str(best.get("url") or ""), allow_generic=True)
        if candidate_domain:
            same_domain = [
                item for item in ranked
                if normalized_domain(str(item.get("url") or ""), allow_generic=True) == candidate_domain
            ]
            if same_domain:
                return same_domain
        if best_score <= 1 and len(results) > 1:
            # Ambiguous snippets are not safe contact evidence. The candidate's
            # website scrape can still enrich the lead without mixing identities.
            return []
        return [item for item in ranked if normalized_domain(str(item.get("url") or ""), allow_generic=True) == best_domain]

    @classmethod
    def _best_email_for_website(
        cls,
        emails: Sequence[str],
        website: str,
        *,
        lead: Optional[Dict[str, Any]] = None,
    ) -> str:
        website_domain = normalized_domain(website, allow_generic=True)
        if lead and (
            lead.get("creator_name")
            or lead.get("instagram_handle")
            or lead.get("tiktok_handle")
        ):
            filtered = cls._creator_emails(
                emails,
                str(lead.get("evidence") or ""),
                name=cls._lead_name(lead),
                socials=cls._lead_socials(lead),
            )
            matching = [email for email in filtered if cls._email_domain(email) == website_domain]
            return matching[0] if matching else (filtered[0] if filtered else "")
        matching = [email for email in emails if cls._email_domain(email) == website_domain]
        return matching[0] if matching else (emails[0] if emails else "")

    @classmethod
    def _creator_emails(
        cls,
        emails: Sequence[str],
        evidence: str = "",
        *,
        name: str = "",
        socials: Optional[Dict[str, Sequence[str]]] = None,
    ) -> List[str]:
        scored: List[Tuple[int, int, str]] = []
        seen: Set[str] = set()
        for raw in emails:
            email = cls._normalize_email(raw)
            if not email or email in seen:
                continue
            if not cls._is_creator_email(email, evidence, name=name, socials=socials):
                continue
            seen.add(email)
            scored.append((cls._creator_email_score(email, evidence, name=name, socials=socials), len(scored), email))
        scored.sort(key=lambda item: (-item[0], item[1]))
        return [email for _, _, email in scored]

    @classmethod
    def _is_creator_email(
        cls,
        email: Any,
        evidence: str = "",
        *,
        name: str = "",
        socials: Optional[Dict[str, Sequence[str]]] = None,
    ) -> bool:
        clean = cls._normalize_email(email)
        if not clean or not EMAIL_RE.fullmatch(clean):
            return False
        local, domain = clean.rsplit("@", 1)
        if cls._is_bad_creator_email_domain(domain):
            return False
        local_key = re.sub(r"[^a-z0-9]", "", local.lower())
        if local.lower() in _BAD_CREATOR_EMAIL_LOCALS or local_key in _BAD_CREATOR_EMAIL_LOCALS:
            return False
        if local_key.startswith(("donotreply", "noreply", "noresponse")):
            return False
        window = cls._email_context_window(clean, evidence)
        if window and any(term in window for term in _BAD_EMAIL_CONTEXT_TERMS):
            positive = {"business", "collab", "collaboration", "contact", "creator", "email", "partnership", "work with"}
            if not any(term in window for term in positive):
                return False
        return True

    @classmethod
    def _creator_email_score(
        cls,
        email: str,
        evidence: str = "",
        *,
        name: str = "",
        socials: Optional[Dict[str, Sequence[str]]] = None,
    ) -> int:
        local, domain = email.rsplit("@", 1)
        local_key = re.sub(r"[^a-z0-9]", "", local.lower())
        domain_key = re.sub(r"[^a-z0-9]", "", domain.lower())
        score = 0
        if domain in _PERSONAL_EMAIL_DOMAINS:
            score += 12
        if local.lower() in _CREATOR_EMAIL_HINT_LOCALS or local_key in _CREATOR_EMAIL_HINT_LOCALS:
            score += 20
        elif any(len(hint) >= 3 and hint in local_key for hint in _CREATOR_EMAIL_HINT_LOCALS):
            score += 8

        handle_tokens = cls._social_identity_tokens(socials)
        for token in handle_tokens:
            if token and (token in local_key or local_key in token):
                score += 30
                break

        for token in cls._name_identity_tokens(name):
            if token in local_key:
                score += 8
            elif token in domain_key:
                score += 4

        window = cls._email_context_window(email, evidence)
        if window:
            if any(term in window for term in ("business", "collab", "collaboration", "contact", "email", "partnership", "work with")):
                score += 10
            if any(term in window for term in _BAD_EMAIL_CONTEXT_TERMS):
                score -= 20
        return score

    @staticmethod
    def _normalize_email(value: Any) -> str:
        match = EMAIL_RE.search(str(value or "").lower())
        if not match:
            return ""
        return match.group(0).strip(".,;:()[]{}").lower()

    @staticmethod
    def _email_domain(email: str) -> str:
        if "@" not in str(email):
            return ""
        return str(email).rsplit("@", 1)[-1].lower().removeprefix("www.")

    @classmethod
    def _is_bad_creator_email_domain(cls, domain: str) -> bool:
        clean = str(domain or "").lower().removeprefix("www.")
        return any(clean == bad or clean.endswith(f".{bad}") for bad in _BAD_CREATOR_EMAIL_DOMAINS)

    @staticmethod
    def _email_context_window(email: str, evidence: str) -> str:
        haystack = str(evidence or "").lower()
        needle = str(email or "").lower()
        index = haystack.find(needle)
        if index < 0:
            return ""
        start = max(0, index - 90)
        end = min(len(haystack), index + len(needle) + 90)
        return haystack[start:end]

    @classmethod
    def _lead_socials(
        cls,
        lead: Dict[str, Any],
        discovered: Optional[Dict[str, Sequence[str]]] = None,
    ) -> Dict[str, List[str]]:
        output: Dict[str, List[str]] = {"instagram": [], "tiktok": []}
        discovered = discovered or {}
        for platform in ("instagram", "tiktok"):
            values: List[Any] = [lead.get(f"{platform}_handle"), lead.get(f"{platform}_url")]
            values.extend(discovered.get(platform) or [])
            for value in values:
                handle = normalize_social_handle(str(value or ""), platform)
                if handle and handle not in output[platform]:
                    output[platform].append(handle)
        return output

    @staticmethod
    def _name_identity_tokens(name: str) -> List[str]:
        return [
            token
            for token in re.findall(r"[a-z0-9]+", str(name or "").lower())
            if len(token) >= 3 and token not in {"and", "the", "dog", "dogs", "pet", "pets", "mom", "dad", "creator", "influencer"}
        ]

    @staticmethod
    def _social_identity_tokens(socials: Optional[Dict[str, Sequence[str]]]) -> List[str]:
        tokens: List[str] = []
        for values in (socials or {}).values():
            for value in values or []:
                clean = re.sub(r"[^a-z0-9]", "", str(value or "").lower().lstrip("@"))
                if len(clean) >= 3 and clean not in tokens:
                    tokens.append(clean)
        return tokens

    async def _scrape_candidate_website(self, website: str) -> Dict[str, Any]:
        empty = {"emails": [], "phones": [], "socials": {"instagram": [], "tiktok": []}, "urls": [], "text": ""}
        if not await self._safe_public_url(website):
            return empty
        headers = {"User-Agent": "Mozilla/5.0 (compatible; VentoLeadFinder/3.0; +https://comfortpod.com/)"}
        timeout = httpx.Timeout(self.scrape_timeout, connect=min(5.0, self.scrape_timeout))
        texts: List[str] = []
        fetched_urls: List[str] = []
        try:
            async with httpx.AsyncClient(timeout=timeout, follow_redirects=True, headers=headers) as client:
                first = await self._fetch_html(client, website)
                if first:
                    final_url, content = first
                    fetched_urls.append(final_url)
                    texts.append(content)
                    contact_urls = []
                    for href in _CONTACT_LINK_RE.findall(content):
                        target = urljoin(final_url, href)
                        if normalized_domain(target, allow_generic=True) == normalized_domain(final_url, allow_generic=True):
                            contact_urls.append(target)
                    if not contact_urls:
                        contact_urls = [urljoin(final_url.rstrip("/") + "/", "contact")]
                    for target in list(dict.fromkeys(contact_urls))[:1]:
                        if target == final_url or not await self._safe_public_url(target):
                            continue
                        page = await self._fetch_html(client, target)
                        if page:
                            fetched_urls.append(page[0])
                            texts.append(page[1])
        except Exception as exc:
            self._stats["scrape_failures"] += 1
            self.logger.info("vento_site_scrape_failed", extra={"website": website, "error": str(exc)})
        combined = " ".join(texts)
        visible = _WHITESPACE_RE.sub(" ", _HTML_TAG_RE.sub(" ", _SCRIPT_STYLE_RE.sub(" ", combined)))
        mailto_emails = extract_emails(" ".join(_MAILTO_RE.findall(combined)))
        visible_emails = extract_emails(visible)
        tel_phones = extract_phones(" ".join(_TEL_RE.findall(combined)))
        json_phones = extract_phones(" ".join(_JSON_PHONE_RE.findall(combined)))
        visible_phones = extract_phones(visible)
        return {
            "emails": list(dict.fromkeys(mailto_emails + visible_emails)),
            "phones": list(dict.fromkeys(tel_phones + json_phones + visible_phones)),
            "socials": extract_social_handles(combined),
            "urls": fetched_urls,
            "text": combined[:10_000],
        }

    async def _fetch_html(self, client: httpx.AsyncClient, url: str) -> Optional[Tuple[str, str]]:
        self._stats["scrape_requests"] += 1
        try:
            response = await client.get(url)
            response.raise_for_status()
            if not await self._safe_public_url(str(response.url)):
                return None
            content_type = response.headers.get("content-type", "").lower()
            if content_type and not any(item in content_type for item in ("html", "text", "xhtml")):
                return None
            return str(response.url), response.text[:1_000_000]
        except Exception:
            self._stats["scrape_failures"] += 1
            return None

    async def _safe_public_url(self, url: str) -> bool:
        parsed = urlparse(str(url or "").strip())
        if parsed.scheme not in {"http", "https"} or not parsed.hostname:
            return False
        host = parsed.hostname.lower()
        if host in {"localhost", "localhost.localdomain"} or host.endswith(".local"):
            return False
        try:
            addresses = await asyncio.to_thread(socket.getaddrinfo, host, parsed.port or (443 if parsed.scheme == "https" else 80))
            for address in addresses:
                ip = ipaddress.ip_address(address[4][0])
                if not ip.is_global:
                    return False
        except (OSError, ValueError):
            return False
        return True

    async def _score_leads(self, leads: List[Dict[str, Any]], category: str, location: str) -> List[Dict[str, Any]]:
        for lead in leads:
            lead["relevance_score"] = self._heuristic_score(lead, category, location)
            lead["scoring_method"] = "heuristic"
        if self.gemini is None or self.gemini_budget <= 0:
            return leads

        semaphore = asyncio.Semaphore(4)

        async def score(lead: Dict[str, Any]) -> None:
            async with semaphore:
                result = await self._gemini_score(lead, category, location)
                if result is not None:
                    lead["relevance_score"] = result[0]
                    lead["scoring_method"] = "gemini"
                    lead["scoring_model"] = result[1]

        await asyncio.gather(*(score(lead) for lead in leads[: self.gemini_budget]))
        return leads

    async def _gemini_score(self, lead: Dict[str, Any], category: str, location: str) -> Optional[Tuple[float, str]]:
        self._stats["gemini_calls"] += 1
        prompt = json.dumps(
            {
                "task": "Score this Vento pet-industry lead using only supplied evidence.",
                "criteria": {"niche_match": "0-4", "has_contact_info": "0-3", "location_match": "0-3"},
                "target": {"category": category, "location": location},
                "lead": {
                    "name": self._lead_name(lead),
                    "email": lead.get("email", ""),
                    "phone": lead.get("phone", ""),
                    "website": lead.get("website", ""),
                    "instagram": lead.get("instagram_handle", ""),
                    "tiktok": lead.get("tiktok_handle", ""),
                    "location": lead.get("location", ""),
                    "evidence": str(lead.get("evidence") or "")[:6_000],
                },
                "output": {"relevance_score": "number from 0 to 10"},
            },
            ensure_ascii=False,
        )
        last_error = ""
        for model in self.gemini_fallback_models:
            self._stats["gemini_api_attempts"] += 1
            try:
                response = await self.gemini.aio.models.generate_content(
                    model=model,
                    contents=prompt,
                    config=genai_types.GenerateContentConfig(
                        temperature=0,
                        max_output_tokens=80,
                        response_mime_type="application/json",
                        system_instruction="Use only evidence provided. Never invent contact or location facts. Return JSON only.",
                    ),
                )
                parsed = json.loads((getattr(response, "text", "") or "{}").strip())
                score = max(0.0, min(10.0, float(parsed.get("relevance_score"))))
                return score, model
            except Exception as exc:
                last_error = str(exc)
                # Authentication, quota and rate-limit failures apply to every
                # model, so another attempt would only add latency/noise.
                lowered = last_error.lower()
                if any(marker in lowered for marker in ("401", "403", "429", "api key", "quota", "rate limit")):
                    break
        self._stats["gemini_failures"] += 1
        self.logger.info("vento_gemini_score_failed", extra={"error": last_error})
        return None

    def _heuristic_score(self, lead: Dict[str, Any], category: str, location: str) -> float:
        evidence = str(lead.get("evidence") or "").lower()
        pet_hits = sum(term in evidence for term in _PET_TERMS)
        niche_points = min(4.0, 1.0 + pet_hits)
        contact_points = min(3.0, float(self._contact_count(lead, category)))
        if self._location_matches(f"{lead.get('location') or ''} {evidence}", location):
            location_points = 3.0
        elif lead.get("location"):
            location_points = 1.0
        else:
            location_points = 0.0
        return round(min(10.0, niche_points + contact_points + location_points), 1)

    def _quality_gate(self, leads: Iterable[Dict[str, Any]], category: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        usable: List[Dict[str, Any]] = []
        rejected: List[Dict[str, Any]] = []
        for raw in leads:
            lead = dict(raw)
            reason = ""
            if not self._lead_name(lead):
                reason = "missing_name"
            elif self._looks_like_placeholder_name(self._lead_name(lead)):
                reason = "placeholder_name"
            elif self._contact_count(lead, category) < 1:
                reason = "missing_contact"
            elif self._clearly_not_pet_related(lead):
                reason = "not_pet_related"
            elif category == "dog_parent_influencers":
                passes_followers, follower_reason = self._meets_min_social_followers(lead)
                if not passes_followers:
                    reason = follower_reason
            if reason:
                lead["usable"] = False
                lead["rejected"] = True
                lead["rejection_reason"] = reason
                rejected.append(lead)
            else:
                lead["usable"] = True
                lead["rejected"] = False
                lead["status"] = "New"
                lead.setdefault("notes", "")
                usable.append(lead)
        return usable, rejected

    def _dedup(self, leads: Iterable[Dict[str, Any]], old_keys: Set[str]) -> Tuple[List[Dict[str, Any]], int]:
        seen = set(old_keys)
        output: List[Dict[str, Any]] = []
        duplicates = 0
        for lead in leads:
            keys = self._dedup_keys(lead)
            if keys and keys.intersection(seen):
                duplicates += 1
                continue
            output.append(dict(lead))
            seen.update(keys)
        return output, duplicates

    def _extract_dedup_keys_from_csv(self, csv_text: str) -> Set[str]:
        text = str(csv_text or "").lstrip("\ufeff")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return self._extract_dedup_keys_from_rows(list(csv.DictReader(io.StringIO(text), dialect=dialect)))

    def _extract_dedup_keys_from_rows(self, rows: Iterable[Dict[str, Any]]) -> Set[str]:
        keys: Set[str] = set()
        for row in rows:
            normalized = {self._normalize_header(key): value for key, value in dict(row).items() if key is not None}
            lead = {
                "creator_name": normalized.get("creatorname") or normalized.get("name") or normalized.get("businessname") or "",
                "email": normalized.get("email") or normalized.get("contactemail") or "",
                "instagram_handle": normalized.get("instagramhandle") or normalized.get("instagram") or "",
                "website": normalized.get("website") or normalized.get("companywebsite") or "",
            }
            keys.update(self._dedup_keys(lead))
        return keys

    def _read_old_lead_rows(self, file_path: str) -> List[Dict[str, Any]]:
        path = Path(file_path)
        if path.suffix.lower() == ".csv":
            raw = path.read_bytes()
            for encoding in ("utf-8-sig", "utf-8", "cp1252"):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                raise ValueError("Unable to decode CSV file")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
        elif path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                worksheet = workbook["Vento Influencers"] if "Vento Influencers" in workbook.sheetnames else workbook.active
                iterator = worksheet.iter_rows(values_only=True)
                headers = [str(value or "").strip() for value in next(iterator, ())]
                rows = [dict(zip(headers, values)) for values in iterator if any(value not in (None, "") for value in values)]
            finally:
                workbook.close()
        else:
            raise ValueError("Only .csv and .xlsx files are supported")
        if len(rows) > self.max_upload_rows:
            raise ValueError(f"Import contains more than {self.max_upload_rows} data rows")
        return rows

    def _parse_file(self, file_path: str) -> List[Dict[str, Any]]:
        """Compatibility parser for older Vento import tests/tools.

        Plan 3 uses uploaded CSV/XLSX mainly for dedup history, but keeping this
        mapper is useful for local imports and does not affect discovery runs.
        """
        path = Path(file_path)
        if path.suffix.lower() == ".xls":
            raise ValueError("Legacy .xls files are not supported; upload .xlsx or .csv")
        if path.suffix.lower() == ".xlsx":
            try:
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    worksheet = workbook.active
                    rows = list(worksheet.iter_rows(values_only=True))
                finally:
                    workbook.close()
            except Exception as exc:
                raise ValueError("Unable to parse XLSX file") from exc
            if not rows:
                return []
            headers = [str(value or "").strip() for value in rows[0]]
            return [self._map_import_row(dict(zip(headers, values))) for values in rows[1:] if any(value not in (None, "") for value in values)]
        if path.suffix.lower() == ".csv":
            return self._parse_raw_text(path.read_text(encoding="utf-8-sig"))
        raise ValueError("Only .csv and .xlsx files are supported")

    def _parse_raw_text(self, raw_text: str) -> List[Dict[str, Any]]:
        text = str(raw_text or "").strip()
        if not text:
            return []
        if text.startswith("[") or text.startswith("{"):
            data = json.loads(text)
            items = data if isinstance(data, list) else [data]
            return [self._map_import_row(item) for item in items if isinstance(item, dict)]
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [self._map_import_row(row) for row in csv.DictReader(io.StringIO(text), dialect=dialect)]

    @classmethod
    def _map_import_row(cls, row: Dict[str, Any]) -> Dict[str, Any]:
        normalized = {cls._normalize_header(key): value for key, value in dict(row or {}).items() if key is not None}
        instagram = normalized.get("instagramurl") or normalized.get("instagram") or normalized.get("ig") or ""
        tiktok = normalized.get("tiktokurl") or normalized.get("tiktok") or normalized.get("tt") or ""
        return {
            "creator_name": _WHITESPACE_RE.sub(
                " ",
                str(
                    normalized.get("fullname")
                    or normalized.get("creator")
                    or normalized.get("name")
                    or normalized.get("creatorname")
                    or ""
                ),
            ).strip(),
            "email": str(normalized.get("businessemail") or normalized.get("email") or normalized.get("contactemail") or "").strip().lower(),
            "instagram_handle": cls._format_handle(normalize_social_handle(str(instagram), "instagram")),
            "tiktok_handle": cls._format_handle(normalize_social_handle(str(tiktok), "tiktok")),
            "follower_count": normalized.get("followers") or normalized.get("followercount") or "",
            "location": str(normalized.get("citystate") or normalized.get("location") or "").strip(),
            "niche": str(normalized.get("niche") or "").strip(),
        }

    @staticmethod
    def _count_csv_rows(csv_text: str) -> int:
        return max(0, sum(1 for row in csv.reader(io.StringIO(str(csv_text or ""))) if row) - 1)

    @classmethod
    def _dedup_keys(cls, lead: Dict[str, Any]) -> Set[str]:
        keys: Set[str] = set()
        name = cls._normalize_name(cls._lead_name(lead))
        email = str(lead.get("email") or "").strip().lower()
        instagram = str(lead.get("instagram_handle") or "").strip().lower().lstrip("@")
        domain = normalized_domain(str(lead.get("website") or ""))
        if name:
            keys.add(f"name:{name}")
        if EMAIL_RE.fullmatch(email):
            keys.add(f"email:{email}")
        if instagram:
            keys.add(f"instagram:{instagram}")
        if domain:
            keys.add(f"domain:{domain}")
        return keys

    @staticmethod
    def _normalize_header(value: Any) -> str:
        return re.sub(r"[^a-z0-9]", "", str(value or "").lower())

    @staticmethod
    def _normalize_name(value: Any) -> str:
        return re.sub(r"[^a-z0-9]", "", str(value or "").lower())

    @classmethod
    def _looks_like_placeholder_name(cls, value: Any) -> bool:
        text = _WHITESPACE_RE.sub(" ", str(value or "")).strip().lower()
        key = cls._normalize_name(text)
        if not key or key in _PLACEHOLDER_NAME_KEYS:
            return True
        tokens = set(re.findall(r"[a-z0-9]+", text))
        return "username" in tokens or bool(tokens and tokens.issubset({"instagram", "tiktok", "user", "profile", "account"}))

    @staticmethod
    def _name_field(category: str) -> str:
        return "creator_name" if category == "dog_parent_influencers" else "business_name"

    @staticmethod
    def _lead_name(lead: Dict[str, Any]) -> str:
        return _WHITESPACE_RE.sub(" ", str(lead.get("creator_name") or lead.get("business_name") or lead.get("name") or "")).strip()

    @staticmethod
    def _format_handle(value: Any) -> str:
        clean = str(value or "").strip().lstrip("@")
        return f"@{clean}" if clean else ""

    @staticmethod
    def _social_url(handle: Any, platform: str) -> str:
        clean = str(handle or "").strip().lstrip("@")
        if not clean:
            return ""
        return f"https://www.instagram.com/{clean}/" if platform == "instagram" else f"https://www.tiktok.com/@{clean}"

    @staticmethod
    def _best_website(websites: Sequence[str], name: str) -> str:
        if not websites:
            return ""
        name_tokens = [token for token in re.findall(r"[a-z0-9]+", name.lower()) if len(token) > 2]
        return max(websites, key=lambda url: sum(token in normalized_domain(url, allow_generic=True) for token in name_tokens))

    @classmethod
    def _location_matches(cls, evidence: str, location: str) -> bool:
        evidence_lower = str(evidence or "").lower()
        parts = [part.strip().lower() for part in re.split(r"[,/]", location) if len(part.strip()) >= 2]
        return bool(parts) and any(part in evidence_lower for part in parts)

    @classmethod
    def _contact_count(cls, lead: Dict[str, Any], category: str) -> int:
        fields = ("email", "instagram_handle", "tiktok_handle") if category == "dog_parent_influencers" else ("email", "phone", "website")
        return sum(bool(str(lead.get(field) or "").strip()) for field in fields)

    def _meets_min_social_followers(self, lead: Dict[str, Any]) -> Tuple[bool, str]:
        if self.min_social_followers <= 0:
            return True, ""
        count = parse_follower_count(lead.get("follower_count"))
        if count is None:
            return (True, "") if not self.require_social_followers else (False, "social_followers_unknown")
        if count < self.min_social_followers:
            return False, f"social_followers_below_{self.min_social_followers}"
        if (
            self.max_social_followers > 0
            and count > self.max_social_followers
            and not self.allow_above_max_social_followers
        ):
            return False, f"social_followers_above_{self.max_social_followers}"
        return True, ""

    def _follower_sort_bucket(self, lead: Dict[str, Any], category: str) -> int:
        if category != "dog_parent_influencers" or self.min_social_followers <= 0:
            return 0
        count = parse_follower_count(lead.get("follower_count"))
        if count is None:
            return 3
        if count < self.min_social_followers:
            return 4
        if self.max_social_followers > 0 and count > self.max_social_followers:
            return 2
        return 0

    @classmethod
    def _extract_follower_count(cls, text: str) -> Optional[int]:
        counts: List[int] = []
        raw = str(text or "")
        for pattern in (_FOLLOWER_COUNT_RE, _FOLLOWER_COUNT_REVERSE):
            for match in pattern.finditer(raw[:500_000]):
                count = parse_follower_count(match.group("count"))
                if count is not None:
                    counts.append(count)
        return max(counts) if counts else None

    def _record_follower_evidence(
        self,
        lead: Dict[str, Any],
        text: str,
        *,
        source_url: str = "",
        status: str = "search_observed",
    ) -> bool:
        count = self._extract_follower_count(text)
        if count is None:
            return False
        return self._set_follower_count(lead, count, source_url=source_url, status=status)

    def _set_follower_count(
        self,
        lead: Dict[str, Any],
        count: int,
        *,
        source_url: str = "",
        status: str = "search_observed",
    ) -> bool:
        current = parse_follower_count(lead.get("follower_count"))
        if current is not None and current >= count:
            return False
        lead["follower_count"] = count
        lead["follower_count_status"] = status
        lead["follower_status"] = status
        lead["follower_source_url"] = source_url
        lead["follower_observed_at"] = datetime.now(timezone.utc).isoformat()
        if str(status).startswith("instagram"):
            lead["instagram_follower_count"] = count
            lead["follower_source_platform"] = "instagram"
        elif str(status).startswith("tiktok"):
            lead["tiktok_follower_count"] = count
            lead["follower_source_platform"] = "tiktok"
        self._stats["social_followers_recovered"] += 1
        return True

    @classmethod
    def _candidate_priority(cls, lead: Dict[str, Any], category: str) -> int:
        evidence = str(lead.get("evidence") or "").lower()
        return cls._contact_count(lead, category) * 4 + sum(term in evidence for term in _PET_TERMS)

    @staticmethod
    def _clearly_not_pet_related(lead: Dict[str, Any]) -> bool:
        evidence = str(lead.get("evidence") or "").lower()
        has_pet = any(term in evidence for term in _PET_TERMS)
        return not has_pet and any(term in evidence for term in _CLEARLY_UNRELATED_TERMS)
