from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from core.services.vento.excel_merge import (
    append_leads_to_csv,
    append_leads_to_workbook,
    inspect_vento_csv,
    inspect_vento_workbook,
)


def test_old_workbook_rows_stay_unchanged_and_new_leads_append(tmp_path):
    source = tmp_path / "old_vento.xlsx"
    output = tmp_path / "updated_vento.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Vento Influencers"
    sheet.append(["Creator ID", "Creator Name", "Email", "Instagram Handle", "Tags", "Custom Column"])
    sheet.append(["old-1", "Old Creator", "old@example.com", "@oldcreator", "Seattle, Pet Creator", "Keep exactly"])
    sheet["A2"].fill = PatternFill(fill_type="solid", fgColor="00FF00")
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "Original second sheet"
    notes["B1"] = "=1+1"
    workbook.save(source)

    sheet_name, _, old_rows, keys = inspect_vento_workbook(str(source))
    assert sheet_name == "Vento Influencers"
    assert old_rows == 1
    assert "instagram:oldcreator" in keys

    stats = append_leads_to_workbook(
        str(source),
        str(output),
        [
            {
                "creator_id": "new-2",
                "creator_name": "New Creator",
                "email": "new@example.com",
                "instagram_handle": "@newcreator",
                "tags": ["Portland", "UGC Creator"],
                "quality_level": "A",
            }
        ],
    )
    assert stats["old_rows"] == 1
    assert stats["new_rows"] == 1

    merged = load_workbook(output, data_only=False)
    try:
        merged_sheet = merged["Vento Influencers"]
        assert [merged_sheet.cell(2, column).value for column in range(1, 7)] == [
            "old-1", "Old Creator", "old@example.com", "@oldcreator", "Seattle, Pet Creator", "Keep exactly"
        ]
        assert merged_sheet["A2"].fill.fgColor.rgb == "0000FF00"
        assert merged_sheet["A3"].value == "new-2"
        assert merged_sheet["B3"].value == "New Creator"
        assert merged_sheet["E3"].value == "Portland, UGC Creator"
        assert merged_sheet["F3"].value in (None, "")
        assert merged["Notes"]["A1"].value == "Original second sheet"
        assert merged["Notes"]["B1"].value == "=1+1"
    finally:
        merged.close()


def test_old_csv_text_is_preserved_and_new_lead_is_appended(tmp_path):
    source = tmp_path / "old_creators.csv"
    output = tmp_path / "old_creators_updated.csv"
    original = 'Name,Email,Instagram,Location,Niche\r\n"Old, Creator",old@example.com,@old,Seattle,Pet Creator'
    source.write_bytes(original.encode("utf-8"))

    _, _, old_rows, keys = inspect_vento_csv(str(source))
    assert old_rows == 1
    assert "instagram:old" in keys
    stats = append_leads_to_csv(
        str(source),
        str(output),
        [
            {
                "creator_name": "New Creator",
                "email": "new@example.com",
                "instagram_handle": "@new",
                "location": "Portland, OR, US",
                "niche": "UGC Creator",
            }
        ],
    )
    merged_bytes = output.read_bytes()
    assert merged_bytes.startswith(original.encode("utf-8"))
    merged_text = merged_bytes.decode("utf-8")
    assert "New Creator,new@example.com,@new" in merged_text
    assert stats == {"sheet_name": "CSV", "old_rows": 1, "new_rows": 1, "total_rows": 2}
